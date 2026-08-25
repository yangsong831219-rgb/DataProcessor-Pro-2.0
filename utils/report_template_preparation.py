"""报告模板检测、标准化与缓存。

用户选择的源文件永不原地修改。所有可接受模板都会复制到指纹缓存中，并写入
``template_manifest.json``。报告构建器只消费缓存副本和该清单中的稳定契约。
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import asdict, dataclass
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Literal, cast

from utils.report_template_validation import validate_report_template


TemplateKind = Literal["word", "ppt"]
TemplateStatus = Literal["direct", "normalized", "rejected"]

_SCHEMA_VERSION = 7
_MANIFEST_NAME = "template_manifest.json"


@dataclass(frozen=True)
class TemplatePreparationResult:
    """一次模板准入的完整、可序列化结果。"""

    report_type: TemplateKind
    status: TemplateStatus
    source_path: str
    usable_path: str | None
    fingerprint: str
    score: int
    reasons: tuple[str, ...]
    warnings: tuple[str, ...]
    capabilities: dict[str, Any]
    profile_path: str | None
    cache_hit: bool = False

    @property
    def accepted(self) -> bool:
        return self.status != "rejected" and bool(self.usable_path)


def prepare_report_template(
    template_path: str,
    report_type: str,
    *,
    cache_root: str | Path | None = None,
) -> TemplatePreparationResult:
    """检测并准备 Word/PPT 模板，返回可直接交给报告构建器的缓存路径。"""
    kind: TemplateKind = "ppt" if report_type == "ppt" else "word"
    source = Path(template_path).expanduser()
    if not source.is_file():
        return _rejected(kind, source, "模板文件不存在或当前不可访问。")

    validation_error = validate_report_template(str(source), kind)
    if validation_error:
        fingerprint = _sha256(source)
        return _cache_rejection(
            kind,
            source,
            fingerprint,
            validation_error,
            cache_root=cache_root,
        )

    fingerprint = _sha256(source)
    cache_dir = _cache_directory(kind, fingerprint, cache_root)
    manifest_path = cache_dir / _MANIFEST_NAME
    cached = _load_cached_result(
        manifest_path,
        expected_kind=kind,
        expected_fingerprint=fingerprint,
    )
    if cached is not None:
        return cached

    cache_dir.mkdir(parents=True, exist_ok=True)
    if kind == "ppt":
        return _prepare_ppt(source, fingerprint, cache_dir)
    return _prepare_word(source, fingerprint, cache_dir)


def load_prepared_template_profile(template_path: str) -> dict[str, Any] | None:
    """读取模板旁的标准化契约；普通、未准备过的模板返回 ``None``。"""
    path = Path(template_path)
    manifest_path = path.parent / _MANIFEST_NAME
    if not path.is_file() or not manifest_path.is_file():
        return None
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return None
    if payload.get("schema_version") != _SCHEMA_VERSION:
        return None
    if payload.get("usable_file") != path.name:
        return None
    profile = payload.get("profile")
    return cast(dict[str, Any], profile) if isinstance(profile, dict) else None


def _prepare_word(
    source: Path,
    fingerprint: str,
    cache_dir: Path,
) -> TemplatePreparationResult:
    from docx import Document

    doc = Document(str(source))
    body = cast(Any, doc._element).body
    meaningful_body_items = _count_meaningful_word_body_items(body)
    section_count = len(doc.sections)
    added_styles = _ensure_required_word_styles(doc)
    body_requires_normalization = meaningful_body_items > 0 or section_count > 1
    requires_normalization = body_requires_normalization or bool(added_styles)
    status: TemplateStatus = "normalized" if requires_normalization else "direct"
    usable_path = cache_dir / "prepared.docx"

    styles = {str(style.name) for style in doc.styles if style.name}
    required_styles = {
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Caption",
        "List Bullet",
        "List Number",
    }
    missing_styles = sorted(required_styles - styles)
    headers = sum(1 for section in doc.sections if _story_has_content(section.header))
    footers = sum(1 for section in doc.sections if _story_has_content(section.footer))
    capabilities: dict[str, Any] = {
        "section_count": section_count,
        "meaningful_body_items": meaningful_body_items,
        "header_count": headers,
        "footer_count": footers,
        "style_count": len(styles),
        "required_styles_present": not missing_styles,
        "missing_required_styles": missing_styles,
    }

    reasons: list[str] = []
    warnings: list[str] = []
    if body_requires_normalization:
        reasons.append("检测到样例正文、表格或多节结构，已清空样例正文。")
        _normalize_word_document(doc)
    else:
        reasons.append("正文为空且节结构简单，可直接作为报告样式模板。")
        if added_styles:
            _normalize_word_document(doc)
    if headers or footers:
        reasons.append("已保留页眉、页脚及页面设置。")
    if added_styles:
        reasons.append("已补齐报告生成所需样式：" + "、".join(added_styles) + "。")
    if missing_styles:
        warnings.append(
            "模板缺少部分标准样式："
            + "、".join(missing_styles)
            + "；报告生成时将由 Word 自动补齐兼容样式。"
        )

    doc.save(str(usable_path))
    score = max(
        0,
        100
        - (8 if requires_normalization else 0)
        - min(20, len(missing_styles) * 4),
    )
    profile = {
        "schema_version": _SCHEMA_VERSION,
        "report_type": "word",
        "body_contract": "empty-report-body",
        "preserve_template_styles": True,
        "preserve_headers_footers": True,
    }
    return _write_result(
        report_type="word",
        status=status,
        source=source,
        usable_path=usable_path,
        fingerprint=fingerprint,
        score=score,
        reasons=reasons,
        warnings=warnings,
        capabilities=capabilities,
        profile=profile,
    )


def _prepare_ppt(
    source: Path,
    fingerprint: str,
    cache_dir: Path,
) -> TemplatePreparationResult:
    from pptx import Presentation

    presentation = Presentation(str(source))
    slide_width = int(presentation.slide_width or 0)
    slide_height = int(presentation.slide_height or 0)
    layout_profiles = [
        _inspect_ppt_layout(layout, index)
        for index, layout in enumerate(presentation.slide_layouts)
    ]
    title_layout = max(
        layout_profiles,
        key=lambda item: int(item["title_score"]),
        default=None,
    )
    content_layout = max(
        layout_profiles,
        key=lambda item: int(item["content_score"]),
        default=None,
    )
    has_title_layout = bool(title_layout and title_layout["has_title"])
    has_content_layout = bool(
        content_layout
        and content_layout["has_title"]
        and content_layout["has_body"]
    )
    editable_cover = _ppt_cover_is_editable(presentation)
    cover_contract = _ppt_cover_contract(presentation)

    needs_cover = cover_contract in {"generated", "standardized-text"}
    uses_safe_textboxes = not has_content_layout
    requires_normalization = needs_cover or uses_safe_textboxes or not has_title_layout
    status: TemplateStatus = "normalized" if requires_normalization else "direct"
    usable_path = cache_dir / "prepared.pptx"

    content_layout_index = (
        int(content_layout["index"]) if content_layout is not None else 0
    )
    title_style, body_style = _infer_ppt_text_styles(
        presentation,
        content_layout_index=content_layout_index,
    )
    profile: dict[str, Any] = {
        "schema_version": _SCHEMA_VERSION,
        "report_type": "ppt",
        "title_layout_index": (
            int(title_layout["index"]) if title_layout is not None else 0
        ),
        "content_layout_index": (
            content_layout_index
        ),
        "content_contract": (
            "native-placeholders" if has_content_layout else "safe-textboxes"
        ),
        "cover_contract": cover_contract,
        "title_style": title_style,
        "body_style": body_style,
        "safe_geometry": {
            "title": [0.06, 0.05, 0.88, 0.12],
            "body": [0.075, 0.22, 0.84, 0.62],
            "body_with_image": [0.075, 0.22, 0.47, 0.62],
            "image": [0.57, 0.22, 0.37, 0.60],
        },
    }
    capabilities: dict[str, Any] = {
        "slide_count": len(presentation.slides),
        "layout_count": len(layout_profiles),
        "slide_width_inches": round(slide_width / 914400.0, 3),
        "slide_height_inches": round(slide_height / 914400.0, 3),
        "editable_cover": editable_cover,
        "cover_contract": cover_contract,
        "native_title_layout": has_title_layout,
        "native_content_layout": has_content_layout,
        "content_contract": profile["content_contract"],
    }

    reasons: list[str] = []
    warnings: list[str] = []
    if cover_contract == "generated":
        reasons.append("模板没有示例封面，已生成可编辑的标准封面。")
    elif cover_contract == "standardized-text":
        reasons.append(
            "模板封面使用样例文字，已保留背景与装饰并转换为标题、作者和日期槽位。"
        )
    else:
        reasons.append("检测到可编辑封面，生成时将保留并填充。")
    if has_content_layout:
        reasons.append("检测到原生标题和正文占位符，可沿用模板版式。")
    else:
        reasons.append("未检测到可靠正文占位符，已建立安全文本框版式契约。")
        warnings.append(
            "该模板主要依赖样例页对象而非母版占位符；基础标准化会保留母版、"
            "背景和推断字体，但不承诺复制样例页的全部装饰对象。"
        )

    if requires_normalization:
        _normalize_ppt_cover(
            presentation,
            title_layout_index=int(profile["title_layout_index"]),
            title_style=title_style,
            body_style=body_style,
            replace_existing_text=cover_contract == "standardized-text",
        )
    presentation.save(str(usable_path))

    score = 100
    if needs_cover:
        score -= 12
    if not has_title_layout:
        score -= 8
    if uses_safe_textboxes:
        score -= 18
    return _write_result(
        report_type="ppt",
        status=status,
        source=source,
        usable_path=usable_path,
        fingerprint=fingerprint,
        score=max(0, score),
        reasons=reasons,
        warnings=warnings,
        capabilities=capabilities,
        profile=profile,
    )


def _normalize_word_document(doc: Any) -> None:
    """清空样例正文，并把已有页眉/页脚引用合并到最终节。"""
    body = cast(Any, doc._element).body
    final_sect_pr = body.sectPr
    if final_sect_pr is None:
        return

    existing_refs: set[tuple[str, str]] = set()
    for ref in list(final_sect_pr):
        local_name = str(ref.tag).rsplit("}", 1)[-1]
        if local_name in {"headerReference", "footerReference"}:
            ref_type = str(ref.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
            ) or "default")
            existing_refs.add((local_name, ref_type))

    for section in list(doc.sections):
        section_properties = cast(Any, section)._sectPr
        for ref in list(section_properties):
            local_name = str(ref.tag).rsplit("}", 1)[-1]
            if local_name not in {"headerReference", "footerReference"}:
                continue
            ref_type = str(ref.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
            ) or "default")
            key = (local_name, ref_type)
            if key not in existing_refs:
                final_sect_pr.insert(0, deepcopy(ref))
                existing_refs.add(key)

    for child in list(body):
        if child is not final_sect_pr:
            body.remove(child)


def _ensure_required_word_styles(doc: Any) -> list[str]:
    from docx.enum.style import WD_STYLE_TYPE

    styles = doc.styles
    required = (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Caption",
        "List Bullet",
        "List Number",
    )
    existing = {str(style.name) for style in styles if style.name}
    added: list[str] = []
    normal = styles["Normal"] if "Normal" in existing else styles.add_style(
        "Normal",
        WD_STYLE_TYPE.PARAGRAPH,
    )
    if "Normal" not in existing:
        added.append("Normal")
        existing.add("Normal")
    for name in required:
        if name in existing:
            continue
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        added.append(name)
        existing.add(name)
    return added


def _count_meaningful_word_body_items(body: Any) -> int:
    count = 0
    for child in list(body):
        local_name = str(child.tag).rsplit("}", 1)[-1]
        if local_name == "sectPr":
            continue
        text = "".join(str(node.text or "") for node in child.iter() if node.text).strip()
        has_graphic = bool(
            child.xpath(".//w:drawing")
            or child.xpath(".//w:pict")
            or child.xpath(".//w:object")
        )
        if local_name == "tbl" or text or has_graphic:
            count += 1
        elif local_name not in {"p", "bookmarkStart", "bookmarkEnd"}:
            count += 1
    return count


def _story_has_content(story: Any) -> bool:
    if any(paragraph.text.strip() for paragraph in story.paragraphs):
        return True
    return bool(story.tables)


def _inspect_ppt_layout(layout: Any, index: int) -> dict[str, Any]:
    from pptx.enum.shapes import PP_PLACEHOLDER

    title_types = {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.VERTICAL_TITLE,
    }
    body_types = {
        PP_PLACEHOLDER.BODY,
        PP_PLACEHOLDER.OBJECT,
        PP_PLACEHOLDER.VERTICAL_BODY,
        PP_PLACEHOLDER.VERTICAL_OBJECT,
    }
    placeholder_types: set[Any] = set()
    for placeholder in layout.placeholders:
        try:
            placeholder_types.add(placeholder.placeholder_format.type)
        except (AttributeError, ValueError):
            continue
    has_title = bool(placeholder_types & title_types)
    has_body = bool(placeholder_types & body_types)
    name = str(getattr(layout, "name", ""))
    lowered = name.lower()
    title_score = (10 if has_title else 0) + (
        3 if "title" in lowered or "标题" in lowered else 0
    )
    content_score = (10 if has_body else 0) + (5 if has_title else 0)
    if any(token in lowered for token in ("content", "text", "正文", "内容")):
        content_score += 3
    if "title slide" in lowered or "标题幻灯片" in lowered:
        content_score -= 4
    return {
        "index": index,
        "name": name,
        "has_title": has_title,
        "has_body": has_body,
        "title_score": title_score,
        "content_score": content_score,
    }


def _ppt_cover_is_editable(presentation: Any) -> bool:
    if len(presentation.slides) == 0:
        return False
    return any(
        getattr(shape, "has_text_frame", False)
        for shape in presentation.slides[0].shapes
    )


def _ppt_cover_contract(presentation: Any) -> str:
    if len(presentation.slides) == 0:
        return "generated"
    cover = presentation.slides[0]
    for shape in cover.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if "{{Report_Title}}" in shape.text or "{{title}}" in shape.text:
            return "tokenized"
    title_shape = cover.shapes.title
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        return "native-title"
    return "standardized-text"


def _infer_ppt_text_styles(
    presentation: Any,
    *,
    content_layout_index: int,
) -> tuple[dict[str, Any], dict[str, Any]]:
    title_candidates: list[Any] = []
    body_candidates: list[Any] = []
    layouts = list(presentation.slide_layouts)
    target_layout = (
        layouts[content_layout_index]
        if 0 <= content_layout_index < len(layouts)
        else None
    )
    all_slides = list(presentation.slides)
    content_slides = [
        slide
        for slide in all_slides[1:]
        if target_layout is not None and slide.slide_layout == target_layout
    ]
    style_slides = content_slides or all_slides[1:] or all_slides
    slide_height = max(1, int(presentation.slide_height or 1))
    for slide in style_slides:
        title_shape = slide.shapes.title
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title_candidates.append(title_shape)
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            if shape is title_shape:
                continue
            max_size = _shape_max_font_size(shape)
            if (
                int(shape.top) < int(slide_height * 0.28)
                or max_size >= 30.0
            ):
                title_candidates.append(shape)
            else:
                body_candidates.append(shape)
    title_candidates.sort(
        key=lambda shape: (_shape_max_font_size(shape), len(shape.text.strip())),
        reverse=True,
    )
    body_candidates.sort(
        key=lambda shape: (
            len(shape.text.strip()),
            -abs(_shape_max_font_size(shape) - 20.0),
        ),
        reverse=True,
    )
    title_style = _first_shape_style(
        title_candidates,
        default={"font_size_pt": 35.0, "bold": True, "color_rgb": "1890FF"},
    )
    body_style = _dominant_body_style(
        body_candidates,
        default={"font_size_pt": 20.0, "bold": False, "color_rgb": "333333"},
    )
    title_style["font_size_pt"] = _clamp_font_size(
        title_style.get("font_size_pt"),
        minimum=30.0,
        maximum=44.0,
        fallback=35.0,
    )
    body_style["font_size_pt"] = _clamp_font_size(
        body_style.get("font_size_pt"),
        minimum=18.0,
        maximum=24.0,
        fallback=20.0,
    )
    return title_style, body_style


def _shape_max_font_size(shape: Any) -> float:
    sizes: list[float] = []
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.font.size is not None:
            sizes.append(float(paragraph.font.size.pt))
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(float(run.font.size.pt))
    return max(sizes, default=0.0)


def _first_shape_style(
    shapes: list[Any],
    *,
    default: dict[str, Any],
) -> dict[str, Any]:
    for shape in shapes:
        for paragraph in shape.text_frame.paragraphs:
            candidates = list(paragraph.runs)
            if not candidates:
                candidates = [paragraph]
            for item in candidates:
                font = item.font
                style: dict[str, Any] = dict(default)
                if font.size is not None:
                    style["font_size_pt"] = round(float(font.size.pt), 2)
                if font.bold is not None:
                    style["bold"] = bool(font.bold)
                if font.name:
                    style["font_name"] = str(font.name)
                color_kind, color_value = _font_color(font)
                if color_kind == "rgb":
                    style["color_rgb"] = color_value
                    style.pop("theme_color", None)
                elif color_kind == "theme":
                    style.pop("color_rgb", None)
                    style["theme_color"] = color_value
                return style
    return dict(default)


def _dominant_body_style(
    shapes: list[Any],
    *,
    default: dict[str, Any],
) -> dict[str, Any]:
    color_scores: dict[tuple[str, Any], int] = {}
    name_scores: dict[str, int] = {}
    size_scores: dict[float, int] = {}
    bold_scores: dict[bool, int] = {False: 0, True: 0}

    for shape in shapes:
        for paragraph in shape.text_frame.paragraphs:
            items = list(paragraph.runs) or [paragraph]
            for item in items:
                text = str(getattr(item, "text", "") or paragraph.text).strip()
                if not text:
                    continue
                font = item.font
                size = (
                    float(font.size.pt)
                    if font.size is not None
                    else _shape_max_font_size(shape)
                )
                if size and not 10.0 <= size <= 28.0:
                    continue
                weight = max(1, min(40, len(text)))
                if size:
                    rounded_size = round(size, 1)
                    size_scores[rounded_size] = (
                        size_scores.get(rounded_size, 0) + weight
                    )
                if font.name:
                    name = str(font.name)
                    name_scores[name] = name_scores.get(name, 0) + weight
                if font.bold is not None:
                    bold_scores[bool(font.bold)] += weight
                color_kind, color_value = _font_color(font)
                if color_kind:
                    key = (color_kind, color_value)
                    color_scores[key] = color_scores.get(key, 0) + weight

    style = dict(default)
    if size_scores:
        style["font_size_pt"] = max(
            size_scores,
            key=lambda item: size_scores[item],
        )
    if name_scores:
        style["font_name"] = max(
            name_scores,
            key=lambda item: name_scores[item],
        )
    if bold_scores[True] or bold_scores[False]:
        style["bold"] = bold_scores[True] > bold_scores[False]
    if color_scores:
        color_kind, color_value = max(
            color_scores,
            key=lambda item: color_scores[item],
        )
        if color_kind == "rgb":
            style["color_rgb"] = color_value
            style.pop("theme_color", None)
        else:
            style.pop("color_rgb", None)
            style["theme_color"] = color_value
    return style


def _font_color(font: Any) -> tuple[str, Any]:
    try:
        if font.color.rgb is not None:
            return "rgb", str(font.color.rgb)
    except (AttributeError, ValueError):
        pass
    try:
        theme_color = font.color.theme_color
        if theme_color is not None and int(theme_color) > 0:
            return "theme", int(theme_color)
    except (AttributeError, TypeError, ValueError):
        pass
    return "", None


def _clamp_font_size(
    value: Any,
    *,
    minimum: float,
    maximum: float,
    fallback: float,
) -> float:
    if not isinstance(value, (int, float)) or isinstance(value, bool) or value <= 0:
        return fallback
    return min(maximum, max(minimum, float(value)))


def _normalize_ppt_cover(
    presentation: Any,
    *,
    title_layout_index: int,
    title_style: dict[str, Any],
    body_style: dict[str, Any],
    replace_existing_text: bool,
) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu

    if len(presentation.slides) == 0:
        layouts = list(presentation.slide_layouts)
        if not layouts:
            raise ValueError("PPT 模板不含可用版式")
        layout = layouts[min(max(0, title_layout_index), len(layouts) - 1)]
        slide = presentation.slides.add_slide(layout)
    else:
        slide = presentation.slides[0]

    if replace_existing_text:
        title_target = _select_cover_title_shape(slide)
        cover_meta_style = (
            _first_shape_style([title_target], default=title_style)
            if title_target is not None
            else dict(title_style)
        )
        cover_meta_style["font_size_pt"] = _clamp_font_size(
            cover_meta_style.get("font_size_pt"),
            minimum=16.0,
            maximum=20.0,
            fallback=18.0,
        )
        cover_meta_style["bold"] = False
        title_element = (
            cast(Any, title_target)._element
            if title_target is not None
            else None
        )
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if title_element is not None and cast(Any, shape)._element is title_element:
                _replace_shape_text_with_token(shape, "{{Report_Title}}")
            else:
                shape.text_frame.clear()
                shape.text_frame.paragraphs[0].text = ""
        if title_target is not None:
            _add_cover_meta_box(
                slide,
                slide_width=int(presentation.slide_width),
                slide_height=int(presentation.slide_height),
                meta_style=cover_meta_style,
                title_target=title_target,
            )
            return

    if not replace_existing_text and any(
        getattr(shape, "has_text_frame", False)
        for shape in slide.shapes
    ):
        return

    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    title_box = slide.shapes.add_textbox(
        Emu(int(slide_width * 0.1)),
        Emu(int(slide_height * 0.3)),
        Emu(int(slide_width * 0.8)),
        Emu(int(slide_height * 0.2)),
    )
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = "{{Report_Title}}"
    title_paragraph.alignment = PP_ALIGN.CENTER
    _apply_ppt_font_style(title_paragraph.font, title_style)

    _add_cover_meta_box(
        slide,
        slide_width=slide_width,
        slide_height=slide_height,
        meta_style=body_style,
        title_target=title_box,
    )


def _select_cover_title_shape(slide: Any) -> Any | None:
    candidates: list[Any] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
            continue
        text = shape.text.strip()
        lowered = text.lower()
        if re.fullmatch(r"(?:19|20)\d{2}|20xx|\d{1,2}", lowered):
            continue
        if re.search(r"\b(?:date|year|author|company name|university name)\b", lowered):
            continue
        candidates.append(shape)
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda shape: (
            _shape_max_font_size(shape),
            int(shape.width) * int(shape.height),
        ),
    )


def _replace_shape_text_with_token(shape: Any, token: str) -> None:
    first_run = None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if first_run is None:
                first_run = run
            run.text = ""
    if first_run is not None:
        first_run.text = token
    else:
        shape.text_frame.clear()
        shape.text_frame.paragraphs[0].text = token


def _add_cover_meta_box(
    slide: Any,
    *,
    slide_width: int,
    slide_height: int,
    meta_style: dict[str, Any],
    title_target: Any,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu

    margin_x = int(slide_width * 0.04)
    left = max(margin_x, min(int(title_target.left), slide_width - margin_x))
    width = min(
        max(int(title_target.width), int(slide_width * 0.2)),
        slide_width - left - margin_x,
    )
    height = int(slide_height * 0.09)
    gap = int(slide_height * 0.025)
    below_title = max(
        int(title_target.top) + int(title_target.height) + gap,
        int(slide_height * 0.42),
    )
    if below_title + height <= int(slide_height * 0.95):
        top = below_title
    else:
        top = max(
            int(slide_height * 0.05),
            int(title_target.top) - height - gap,
        )

    meta_box = slide.shapes.add_textbox(
        Emu(left),
        Emu(top),
        Emu(width),
        Emu(height),
    )
    meta_box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    meta_paragraph = meta_box.text_frame.paragraphs[0]
    meta_paragraph.text = (
        "{{Report_Author}}\v{{Report_Date}}"
        if width < int(slide_width * 0.45)
        else "{{Report_Author}}  |  {{Report_Date}}"
    )
    title_alignment = title_target.text_frame.paragraphs[0].alignment
    meta_paragraph.alignment = title_alignment or PP_ALIGN.LEFT
    _apply_ppt_font_style(meta_paragraph.font, meta_style)


def _apply_ppt_font_style(font: Any, style: dict[str, Any]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Pt

    size = style.get("font_size_pt")
    if isinstance(size, (int, float)) and size > 0:
        font.size = Pt(float(size))
    bold = style.get("bold")
    if isinstance(bold, bool):
        font.bold = bold
    name = style.get("font_name")
    if isinstance(name, str) and name:
        font.name = name
    color = style.get("color_rgb")
    if isinstance(color, str) and len(color) == 6:
        try:
            font.color.rgb = RGBColor.from_string(color)
        except ValueError:
            pass
    theme_color = style.get("theme_color")
    if isinstance(theme_color, int) and not isinstance(theme_color, bool):
        try:
            font.color.theme_color = MSO_THEME_COLOR(theme_color)
        except ValueError:
            pass


def _cache_rejection(
    report_type: TemplateKind,
    source: Path,
    fingerprint: str,
    reason: str,
    *,
    cache_root: str | Path | None,
) -> TemplatePreparationResult:
    cache_dir = _cache_directory(report_type, fingerprint, cache_root)
    cache_dir.mkdir(parents=True, exist_ok=True)
    return _write_result(
        report_type=report_type,
        status="rejected",
        source=source,
        usable_path=None,
        fingerprint=fingerprint,
        score=0,
        reasons=[reason],
        warnings=[],
        capabilities={},
        profile={},
        cache_dir=cache_dir,
    )


def _rejected(
    report_type: TemplateKind,
    source: Path,
    reason: str,
) -> TemplatePreparationResult:
    return TemplatePreparationResult(
        report_type=report_type,
        status="rejected",
        source_path=str(source),
        usable_path=None,
        fingerprint="",
        score=0,
        reasons=(reason,),
        warnings=(),
        capabilities={},
        profile_path=None,
    )


def _write_result(
    *,
    report_type: TemplateKind,
    status: TemplateStatus,
    source: Path,
    usable_path: Path | None,
    fingerprint: str,
    score: int,
    reasons: list[str],
    warnings: list[str],
    capabilities: dict[str, Any],
    profile: dict[str, Any],
    cache_dir: Path | None = None,
) -> TemplatePreparationResult:
    target_dir = cache_dir or cast(Path, usable_path).parent
    manifest_path = target_dir / _MANIFEST_NAME
    result = TemplatePreparationResult(
        report_type=report_type,
        status=status,
        source_path=str(source.resolve()),
        usable_path=str(usable_path.resolve()) if usable_path is not None else None,
        fingerprint=fingerprint,
        score=score,
        reasons=tuple(reasons),
        warnings=tuple(warnings),
        capabilities=capabilities,
        profile_path=str(manifest_path.resolve()),
    )
    payload = asdict(result)
    payload.update(
        {
            "schema_version": _SCHEMA_VERSION,
            "usable_file": usable_path.name if usable_path is not None else None,
            "profile": profile,
        }
    )
    temporary = manifest_path.with_suffix(".tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary.replace(manifest_path)
    return result


def _load_cached_result(
    manifest_path: Path,
    *,
    expected_kind: TemplateKind,
    expected_fingerprint: str,
) -> TemplatePreparationResult | None:
    if not manifest_path.is_file():
        return None
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return None
    if (
        payload.get("schema_version") != _SCHEMA_VERSION
        or payload.get("report_type") != expected_kind
        or payload.get("fingerprint") != expected_fingerprint
    ):
        return None
    usable_file = payload.get("usable_file")
    usable_path = manifest_path.parent / usable_file if isinstance(usable_file, str) else None
    if payload.get("status") != "rejected" and (
        usable_path is None or not usable_path.is_file()
    ):
        return None
    try:
        return TemplatePreparationResult(
            report_type=expected_kind,
            status=cast(TemplateStatus, payload["status"]),
            source_path=str(payload["source_path"]),
            usable_path=str(usable_path.resolve()) if usable_path is not None else None,
            fingerprint=expected_fingerprint,
            score=int(payload["score"]),
            reasons=tuple(str(item) for item in payload.get("reasons", [])),
            warnings=tuple(str(item) for item in payload.get("warnings", [])),
            capabilities=cast(
                dict[str, Any],
                payload.get("capabilities", {}),
            ),
            profile_path=str(manifest_path.resolve()),
            cache_hit=True,
        )
    except (KeyError, TypeError, ValueError):
        return None


def _cache_directory(
    report_type: TemplateKind,
    fingerprint: str,
    cache_root: str | Path | None,
) -> Path:
    root = Path(cache_root) if cache_root is not None else _default_cache_root()
    return root / report_type / fingerprint[:20]


def _default_cache_root() -> Path:
    local_app_data = os.environ.get("LOCALAPPDATA")
    if local_app_data:
        return Path(local_app_data) / "DataProcessorPro" / "report_template_cache"
    return Path(tempfile.gettempdir()) / "DataProcessorPro" / "report_template_cache"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
