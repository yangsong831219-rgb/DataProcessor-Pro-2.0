"""无模板 PPT 也必须具备正式技术汇报的稳定视觉系统。"""

from __future__ import annotations

from PIL import Image
from pptx import Presentation

from dp_engine.report_builder.models import PPTReport, PPTSlide
from dp_engine.report_builder.ppt_builder import PPTBuilder


def test_default_ppt_uses_formal_theme_and_content_chart(tmp_path) -> None:
    chart = tmp_path / "compare.png"
    Image.new("RGB", (1280, 720), color=(235, 242, 248)).save(chart)
    output = tmp_path / "formal.pptx"
    report = PPTReport(
        title="多源传感器诊断与决策汇报",
        author="DataProcessor Pro",
        slides=[
            PPTSlide(
                slide_title="多源一致性不足，当前判废结论需复核",
                bullet_points=["相关性偏低", "测点映射缺失", "建议补充交叉验证"],
                image_anchor="[INSERT_IMAGE: compare.png]",
            )
        ],
    )

    PPTBuilder().build_ppt_report(
        report,
        "",
        str(output),
        project_dir=str(tmp_path),
    )

    prs = Presentation(str(output))
    assert len(prs.slides) == 2
    assert str(prs.slides[0].background.fill.fore_color.rgb) == "0B1F33"
    content = prs.slides[1]
    assert str(content.background.fill.fore_color.rgb) == "F7F9FC"
    assert sum(1 for shape in content.shapes if shape.shape_type == 13) == 1
    assert any(
        "DATAPROCESSOR PRO" in shape.text
        for shape in content.shapes
        if getattr(shape, "has_text_frame", False)
    )
    title = content.shapes.title
    assert title is not None
    assert title.text == "多源一致性不足，当前判废结论需复核"
    assert title.text_frame.paragraphs[0].font.size.pt >= 35
