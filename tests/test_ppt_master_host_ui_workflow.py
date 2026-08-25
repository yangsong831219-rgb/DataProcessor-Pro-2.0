from __future__ import annotations

from pathlib import Path
from typing import TypeVar

import pytest
from pydantic import BaseModel, ValidationError
from pptx import Presentation
from pptx.util import Inches

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QScrollArea, QSplitter

from dp_engine.ppt_master_host import (
    ColorPalette,
    CommunicationContract,
    DesignContract,
    OutlinePlan,
    OutlineSection,
    PlanningAsset,
    PlanningAssetKind,
    PlanningPhase,
    PlanningRequest,
    PptMasterPlanningWorkflow,
    PptMasterTemplateWorkspaceError,
    PptMasterWorkflowError,
    SlideIntent,
    SlideIntentPlan,
    SlideLayout,
    TemplateMode,
    TypographySpec,
    fingerprint_ppt_master_inputs,
    prepare_ppt_master_template_workspace,
)
from dp_engine.ppt_master_host.authoring import (
    HostAIPptMasterAuthoringAdapter,
)
from dp_engine.report_provider import (
    BUILTIN_PROVIDER_ID,
    BUILTIN_PROVIDER_VERSION,
    PPT_MASTER_PROVIDER_ID,
    PPT_MASTER_PROVIDER_VERSION,
    PptMasterAuthoringAsset,
    PptMasterAuthoringContext,
    PptMasterReportRenderProvider,
    PptMasterSlideAuthoringRequest,
    ReportProviderOption,
    ReportRenderAsset,
    ReportRenderRequest,
)
from ui.report_workbench import ReportWorkbenchWidget
from utils.report_template_preparation import prepare_report_template


ModelT = TypeVar("ModelT", bound=BaseModel)


class ScriptedModel:
    identity = "scripted:batch-3.6.5"

    def __init__(self, outputs: list[BaseModel]) -> None:
        self.outputs = list(outputs)

    def generate_structured(
        self,
        *,
        prompt: str,
        schema: type[ModelT],
        system_prompt: str,
        max_tokens: int,
        max_schema_retries: int = 1,
    ) -> ModelT:
        del prompt, system_prompt, max_tokens, max_schema_retries
        return schema.model_validate(self.outputs.pop(0))


class ScriptedAuthoringModel:
    identity = "scripted:authoring"

    def __init__(self, svg_text: str | list[str]) -> None:
        self.svg_texts = [svg_text] if isinstance(svg_text, str) else list(svg_text)
        self.prompts: list[str] = []

    def generate_structured(
        self,
        *,
        prompt: str,
        schema: type[ModelT],
        system_prompt: str,
        max_tokens: int,
    ) -> ModelT:
        self.prompts.append(prompt)
        del system_prompt, max_tokens
        return schema.model_validate({
            "svg_text": self.svg_texts.pop(0),
            "speaker_notes_markdown": "说明诊断证据及其决策含义。",
        })


def test_authored_svg_payload_rejects_contract_invalid_svg_before_provider_stage() -> None:
    """An invalid SVG must enter the model's per-slide corrective retry path."""
    from dp_engine.ppt_master_host.authoring import _AuthoredSvgPayload

    with pytest.raises(
        ValidationError,
        match="SVG image reference escapes the project pool",
    ):
        _AuthoredSvgPayload.model_validate({
            "svg_text": (
                '<svg xmlns="http://www.w3.org/2000/svg" '
                'viewBox="0 0 1280 720">'
                '<image href="https://invalid.example/chart.png" '
                'x="64" y="120" width="700" height="500"/>'
                "</svg>"
            ),
            "speaker_notes_markdown": "invalid fixture",
        })


def test_ai_client_retries_only_current_authored_slide_after_svg_contract_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """SVG contract errors participate in the existing bounded schema retry."""
    from core.ai_client import AIClient
    from dp_engine.ppt_master_host.authoring import _AuthoredSvgPayload

    client = object.__new__(AIClient)
    object.__setattr__(client, "_initialized", False)
    AIClient.__init__(client, api_key="sk-test")
    responses = iter((
        '{"svg_text":"<svg xmlns=\\"http://www.w3.org/2000/svg\\" '
        'viewBox=\\"0 0 1280 720\\"><image href=\\"https://invalid.example/'
        'chart.png\\" x=\\"64\\" y=\\"120\\" width=\\"700\\" height=\\"500\\"/>'
        '</svg>","speaker_notes_markdown":"invalid"}',
        '{"svg_text":"<svg xmlns=\\"http://www.w3.org/2000/svg\\" '
        'viewBox=\\"0 0 1280 720\\"><rect x=\\"0\\" y=\\"0\\" width=\\"1280\\" '
        'height=\\"720\\" fill=\\"#0D1117\\"/></svg>",'
        '"speaker_notes_markdown":"corrected"}',
    ))
    calls: list[str] = []

    def _generate(**kwargs: object) -> str:
        calls.append(str(kwargs["prompt"]))
        return next(responses)

    monkeypatch.setattr(client, "generate", _generate)

    result = client.generate_structured(
        "author slide 07",
        _AuthoredSvgPayload,
        max_schema_retries=1,
    )

    assert result["speaker_notes_markdown"] == "corrected"
    assert len(calls) == 2
    assert "SVG image reference escapes the project pool" in calls[1]


def test_host_authorer_retries_current_slide_on_exact_asset_reference_mismatch() -> None:
    """A valid SVG using an adjacent slide asset is corrected in Host authoring."""
    asset_ids = (
        "compare_corr_scatter_1",
        "compare_corr_scatter_2",
    )
    project_filenames = (
        "asset_02_compare_corr_scatter_1.png",
        "asset_03_compare_corr_scatter_2.png",
    )
    request = PlanningRequest(
        request_id="slide-13-asset-contract",
        report_title="多源交叉比对",
        objective="用两组相关性图定位异常通道",
        audience="项目技术负责人",
        source_context="第 13 页应精确引用计划中的两张相关性图。",
        requested_slide_count=3,
        assets=tuple(
            PlanningAsset(
                asset_id=asset_id,
                kind=PlanningAssetKind.CHART,
                semantic_label=asset_id,
                summary="相关性证据",
                required=True,
            )
            for asset_id in asset_ids
        ),
    )
    outline = OutlinePlan(
        deck_title=request.report_title,
        narrative_arc="从异常到证据再到决策",
        sections=(OutlineSection(
            section_id="cross_comparison",
            title="相关性验证",
            purpose="定位异常通道",
            key_messages=("两张图必须同页对照",),
            allocated_slides=3,
            candidate_asset_ids=asset_ids,
        ),),
    )
    design = _planning_outputs(template_mode=TemplateMode.FREE_DESIGN)[1]
    slides = SlideIntentPlan(
        deck_title=outline.deck_title,
        slides=(
            SlideIntent(
                slide_id="slide-01",
                section_id="cross_comparison",
                sequence=1,
                title="分析背景",
                message="建立分析语境",
                layout=SlideLayout.TITLE,
                content_points=("异常通道",),
                visual_brief="标题页",
            ),
            SlideIntent(
                slide_id="slide-13",
                section_id="cross_comparison",
                sequence=2,
                title="多源交叉比对：相关性验证（二）",
                message="两组相关性均很弱",
                layout=SlideLayout.COMPARISON,
                content_points=("r=-0.2472", "r=0.1672"),
                asset_ids=asset_ids,
                visual_brief="左右并排两张相关性散点图",
                speaker_notes_intent="说明异常通道判断",
            ),
            SlideIntent(
                slide_id="slide-03",
                section_id="cross_comparison",
                sequence=3,
                title="处置决策",
                message="排查通道映射",
                layout=SlideLayout.CONCLUSION,
                content_points=("复核映射",),
                visual_brief="决策页",
            ),
        ),
    )
    fingerprint = "d" * 64
    workflow = PptMasterPlanningWorkflow(
        ScriptedModel([outline, design, slides]),
        input_fingerprint=fingerprint,
    )
    workflow.generate_outline(request)
    workflow.advance("confirm_outline", current_input_fingerprint=fingerprint)
    workflow.advance("confirm_design", current_input_fingerprint=fingerprint)
    workflow.advance("confirm_plan", current_input_fingerprint=fingerprint)
    assert workflow.snapshot is not None

    context = PptMasterAuthoringContext(
        planning_snapshot=workflow.snapshot,
        render_request_sha256="e" * 64,
        structured_report_json='{"slides":[{},{},{}]}',
        slide_filenames=("P01_slide_01.svg", "P13_slide_13.svg", "P03_slide_03.svg"),
        assets=tuple(
            PptMasterAuthoringAsset(
                asset_id=asset_id,
                project_filename=filename,
                media_type="image/png",
                semantic_label=asset_id,
                target="compare",
                required=True,
                size_bytes=1,
                sha256="f" * 64,
            )
            for asset_id, filename in zip(asset_ids, project_filenames)
        ),
    )
    wrong_svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
        '<rect width="1280" height="720" fill="#F7F9FC"/>'
        '<image href="../images/asset_01_compare_corr_scatter_0.png" '
        'x="64" y="120" width="560" height="420"/></svg>'
    )
    corrected_svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
        '<rect width="1280" height="720" fill="#F7F9FC"/>'
        f'<image href="../images/{project_filenames[0]}" '
        'x="64" y="120" width="540" height="420"/>'
        f'<image href="../images/{project_filenames[1]}" '
        'x="676" y="120" width="540" height="420"/></svg>'
    )
    model = ScriptedAuthoringModel([wrong_svg, corrected_svg])
    adapter = HostAIPptMasterAuthoringAdapter(model)

    authored = adapter.author_slide(
        context,
        PptMasterSlideAuthoringRequest(
            intent=slides.slides[1],
            svg_filename=context.slide_filenames[1],
        ),
    )

    assert len(model.prompts) == 2
    assert "asset_01_compare_corr_scatter_0.png" in model.prompts[1]
    assert project_filenames[0] in model.prompts[1]
    assert project_filenames[1] in model.prompts[1]
    assert "asset_01_compare_corr_scatter_0.png" not in authored.svg_text
    assert all(filename in authored.svg_text for filename in project_filenames)
    assert authored.used_asset_ids == asset_ids


def test_host_authorer_fails_after_bounded_asset_reference_retry() -> None:
    """A repeated wrong image set fails closed with the precise Provider code."""
    workflow = _confirmed_workflow()
    assert workflow.snapshot is not None
    asset = PptMasterAuthoringAsset(
        asset_id="图表.温度",
        project_filename="asset_01_item.png",
        media_type="image/png",
        semantic_label="温度时程曲线",
        target="温度诊断",
        required=True,
        size_bytes=1,
        sha256="f" * 64,
    )
    context = PptMasterAuthoringContext(
        planning_snapshot=workflow.snapshot,
        render_request_sha256="e" * 64,
        structured_report_json='{"slides":[{},{},{}]}',
        slide_filenames=("P01_slide_01.svg", "P02_slide_02.svg", "P03_slide_03.svg"),
        assets=(asset,),
    )
    wrong_svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
        '<rect width="1280" height="720" fill="#F7F9FC"/>'
        '<image href="../images/asset_99_wrong.png" '
        'x="64" y="120" width="700" height="420"/></svg>'
    )
    model = ScriptedAuthoringModel([wrong_svg, wrong_svg])
    adapter = HostAIPptMasterAuthoringAdapter(model)
    slides = workflow.snapshot.slides
    assert slides is not None

    from dp_engine.ppt_master_host import HostAuthoringError

    with pytest.raises(HostAuthoringError) as exc_info:
        adapter.author_slide(
            context,
            PptMasterSlideAuthoringRequest(
                intent=slides.slides[1],
                svg_filename=context.slide_filenames[1],
            ),
        )

    assert exc_info.value.code == "svg_asset_reference_mismatch"
    assert len(model.prompts) == 2


class UnusedRunner:
    def workspace_path(self, workspace_id: str) -> Path:
        raise AssertionError(workspace_id)

    def run(self, *args, **kwargs):
        raise AssertionError((args, kwargs))


def _planning_outputs(
    *,
    template_mode: TemplateMode,
    asset_id: str = "图表.温度",
) -> list[BaseModel]:
    outline = OutlinePlan(
        deck_title="温度补偿诊断结论",
        narrative_arc="从事实证据推进到处置决策",
        sections=(
            OutlineSection(
                section_id="evidence",
                title="关键证据",
                purpose="解释诊断图所揭示的问题",
                key_messages=("温度补偿结果决定可用性",),
                allocated_slides=3,
                candidate_asset_ids=(asset_id,),
            ),
        ),
    )
    design = DesignContract(
        communication=CommunicationContract(
            objective="形成可审核的处置建议",
            audience_success="审核人能依据证据作出决策",
            tone="严谨、克制、结论优先",
            content_divergence="faithful",
        ),
        template_strategy=template_mode,
        visual_style="正式工程技术汇报",
        palette=ColorPalette(
            primary="#164C7E",
            secondary="#527A9E",
            accent="#E07A2D",
            background="#F7F9FC",
            text="#17212B",
        ),
        typography=TypographySpec(
            title_font="Microsoft YaHei",
            body_font="Microsoft YaHei",
            monospace_font="Consolas",
            title_size_px=38,
            body_size_px=22,
            caption_size_px=15,
        ),
        density="balanced",
        image_usage="source_only",
        chart_style="突出关键趋势和阈值",
        structure_mode="flat",
        refine_spec="使用统一网格和安全边距",
        layout_rules=("每页一个结论", "图片与解释同页"),
        accessibility_rules=("正文保持高对比度",),
    )
    slides = SlideIntentPlan(
        deck_title=outline.deck_title,
        slides=(
            SlideIntent(
                slide_id="slide-01",
                section_id="evidence",
                sequence=1,
                title="诊断目标与判断门槛",
                message="先明确验收问题与证据边界",
                layout=SlideLayout.TITLE,
                content_points=("说明对象", "明确阈值"),
                visual_brief="使用克制的标题页",
                speaker_notes_intent="交代报告范围",
            ),
            SlideIntent(
                slide_id="slide-02",
                section_id="evidence",
                sequence=2,
                title="温度证据支持补偿判断",
                message="图表直接支撑补偿有效性判断",
                layout=SlideLayout.DATA_FOCUS,
                content_points=("读取趋势", "解释偏差"),
                asset_ids=(asset_id,),
                visual_brief="左侧图表，右侧结论",
                speaker_notes_intent="解释图片中的证据",
            ),
            SlideIntent(
                slide_id="slide-03",
                section_id="evidence",
                sequence=3,
                title="形成分级处置结论",
                message="按风险等级给出后续动作",
                layout=SlideLayout.CONCLUSION,
                content_points=("立即动作", "复验条件"),
                visual_brief="结论卡片和行动序列",
                speaker_notes_intent="说明处置优先级",
            ),
        ),
    )
    return [outline, design, slides]


def _request(
    *,
    template_mode: TemplateMode,
    template_summary: str = "",
    asset_id: str = "图表.温度",
) -> PlanningRequest:
    return PlanningRequest(
        request_id="batch-365-plan",
        report_title="温度补偿诊断结论",
        objective="形成可审核的处置建议",
        audience="项目技术负责人",
        source_context="诊断记录显示温度补偿前后存在可比较证据。",
        requested_slide_count=3,
        template_mode=template_mode,
        template_summary=template_summary,
        assets=(PlanningAsset(
            asset_id=asset_id,
            kind=PlanningAssetKind.CHART,
            semantic_label="温度时程曲线",
            summary="用于温度补偿判断",
            required=True,
        ),),
    )


def _confirmed_workflow(
    *,
    template_mode: TemplateMode = TemplateMode.FREE_DESIGN,
    template_summary: str = "",
) -> PptMasterPlanningWorkflow:
    fingerprint = "a" * 64
    workflow = PptMasterPlanningWorkflow(
        ScriptedModel(_planning_outputs(template_mode=template_mode)),
        input_fingerprint=fingerprint,
    )
    workflow.generate_outline(_request(
        template_mode=template_mode,
        template_summary=template_summary,
    ))
    workflow.advance("confirm_outline", current_input_fingerprint=fingerprint)
    workflow.advance("confirm_design", current_input_fingerprint=fingerprint)
    workflow.advance("confirm_plan", current_input_fingerprint=fingerprint)
    return workflow


def _make_prepared_template(tmp_path: Path, *, widescreen: bool = True) -> Path:
    tmp_path.mkdir(parents=True, exist_ok=True)
    source = tmp_path / ("wide.pptx" if widescreen else "standard.pptx")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333 if widescreen else 10)
    presentation.slide_height = Inches(7.5)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(str(source))
    result = prepare_report_template(
        str(source),
        "ppt",
        cache_root=tmp_path / "prepared-cache",
    )
    assert result.accepted and result.usable_path
    return Path(result.usable_path)


def test_staged_workflow_requires_each_confirmation_and_input_fingerprint() -> None:
    config = {
        "report_type": "ppt",
        "req_file": "D:/project/requirement.txt",
        "template_file": "",
        "project_files": ["D:/project/data.csv"],
        "_diagnosis_record": {
            "record_id": "record-1",
            "chart_manifest": [{"fig_id": "图表.温度"}],
        },
        "report_provider": {"provider_id": PPT_MASTER_PROVIDER_ID},
    }
    fingerprint = fingerprint_ppt_master_inputs(config)
    workflow = PptMasterPlanningWorkflow(
        ScriptedModel(_planning_outputs(template_mode=TemplateMode.FREE_DESIGN)),
        input_fingerprint=fingerprint,
    )
    workflow.generate_outline(_request(template_mode=TemplateMode.FREE_DESIGN))
    assert workflow.snapshot and workflow.snapshot.phase == PlanningPhase.OUTLINE_READY
    assert workflow.view().next_action == "confirm_outline"

    workflow.advance("confirm_outline", current_input_fingerprint=fingerprint)
    assert workflow.view().next_action == "confirm_design"
    workflow.advance("confirm_design", current_input_fingerprint=fingerprint)
    assert workflow.view().next_action == "confirm_plan"
    workflow.advance("confirm_plan", current_input_fingerprint=fingerprint)
    assert workflow.ready_for_authoring
    report_slides = workflow.structured_report()["slides"]
    assert isinstance(report_slides, list)
    assert len(report_slides) == 3

    changed = dict(config)
    changed["template_file"] = "D:/project/new-template.pptx"
    with pytest.raises(PptMasterWorkflowError, match="输入已改变"):
        workflow.require_current_inputs(fingerprint_ppt_master_inputs(changed))
    assert not workflow.ready_for_authoring


def test_template_workspace_accepts_prepared_169_and_rejects_43(tmp_path: Path) -> None:
    prepared = _make_prepared_template(tmp_path, widescreen=True)
    workspace = prepare_ppt_master_template_workspace(
        prepared,
        workspace_root=tmp_path / "workspaces",
    )
    cached = prepare_ppt_master_template_workspace(
        prepared,
        workspace_root=tmp_path / "workspaces",
    )
    assert cached == workspace
    assert workspace.template_path.is_file()
    assert workspace.template_sha256
    assert "不复制样例对象" in workspace.style_summary

    standard = _make_prepared_template(tmp_path / "standard", widescreen=False)
    with pytest.raises(PptMasterTemplateWorkspaceError) as captured:
        prepare_ppt_master_template_workspace(
            standard,
            workspace_root=tmp_path / "workspaces-43",
        )
    assert captured.value.code == "canvas_unsupported"


def test_template_attestation_and_concrete_authoring_are_path_free(tmp_path: Path) -> None:
    prepared = _make_prepared_template(tmp_path, widescreen=True)
    template_workspace = prepare_ppt_master_template_workspace(
        prepared,
        workspace_root=tmp_path / "workspaces",
    )
    workflow = PptMasterPlanningWorkflow(
        ScriptedModel(_planning_outputs(
            template_mode=TemplateMode.VALIDATED_WORKSPACE,
        )),
        input_fingerprint="b" * 64,
        template_workspace=template_workspace,
    )
    workflow.generate_outline(_request(
        template_mode=TemplateMode.VALIDATED_WORKSPACE,
        template_summary=template_workspace.style_summary,
    ))
    workflow.advance("confirm_outline", current_input_fingerprint="b" * 64)
    workflow.advance("confirm_design", current_input_fingerprint="b" * 64)
    workflow.advance("confirm_plan", current_input_fingerprint="b" * 64)
    assert workflow.snapshot is not None

    image_path = (tmp_path / "temperature.png").resolve()
    image_path.write_bytes(b"png-fixture")
    output = (tmp_path / "output.pptx").resolve()
    output.write_bytes(b"")
    render_asset = ReportRenderAsset(
        host_id="图表.温度",
        source_path=image_path,
        media_type="image/png",
        semantic_label="温度时程曲线",
        target="温度诊断",
    )
    provider = PptMasterReportRenderProvider(
        planning_snapshot=workflow.snapshot,
        runner=UnusedRunner(),
        authoring_adapter=HostAIPptMasterAuthoringAdapter(
            ScriptedAuthoringModel("<svg xmlns='http://www.w3.org/2000/svg' "
                "viewBox='0 0 1280 720'><rect width='1280' height='720' "
                "fill='#F7F9FC'/><image href='../images/asset_01_item.png' "
                "x='64' y='120' width='700' height='500'/></svg>")
        ),
        template_workspace=template_workspace,
    )
    render_request = ReportRenderRequest(
        report_type="ppt",
        structured_report=workflow.structured_report(),
        template_path=str(prepared),
        output_path=str(output),
        assets=(render_asset,),
        required_figure_ids=("图表.温度",),
    )
    prepared_request = provider._prepare_request(  # pyright: ignore[reportPrivateUsage]
        workflow.snapshot,
        render_request,
    )
    assert prepared_request.context.template_style is not None
    assert str(tmp_path) not in prepared_request.context.model_dump_json()

    adapter = provider._authoring_adapter  # pyright: ignore[reportPrivateUsage]
    project_spec = adapter.create_project_spec(prepared_request.context)
    assert "ppt-master-schema: design-spec/v1" in project_spec.design_spec_markdown
    assert "[fill" not in project_spec.design_spec_markdown.lower()
    assert "- caption: 15" in project_spec.spec_lock_markdown
    slide = workflow.snapshot.slides.slides[1]  # type: ignore[union-attr]
    authored = adapter.author_slide(
        prepared_request.context,
        PptMasterSlideAuthoringRequest(
            intent=slide,
            svg_filename=prepared_request.context.slide_filenames[1],
        ),
    )
    assert authored.used_asset_ids == ("图表.温度",)


def test_workbench_exposes_host_confirmation_and_explicit_fallback(qapp) -> None:
    widget = ReportWorkbenchWidget()
    widget.set_report_provider_options((
        ReportProviderOption(
            provider_id=BUILTIN_PROVIDER_ID,
            provider_version=BUILTIN_PROVIDER_VERSION,
            display_name="内置标准生成器",
            is_builtin=True,
        ),
        ReportProviderOption(
            provider_id=PPT_MASTER_PROVIDER_ID,
            provider_version=PPT_MASTER_PROVIDER_VERSION,
            display_name="PPT Master 专业演示生成器",
        ),
    ))
    for index in range(widget.report_provider_combo.count()):
        data = widget.report_provider_combo.itemData(index) or {}
        if data.get("provider_id") == PPT_MASTER_PROVIDER_ID:
            widget.report_provider_combo.setCurrentIndex(index)
            break
    assert widget.ppt_master_group.isVisibleTo(widget)
    assert widget.outline_editor.isReadOnly()

    widget.set_diagnosis_record({"timestamp": "2026-08-06"})
    widget.set_ppt_master_workflow_view(
        status_text="规划已确认",
        preview_markdown="# 已确认方案",
        ready_for_authoring=True,
    )
    assert widget.full_report_btn.isEnabled()
    widget.ppt_master_fallback_btn.click()
    assert widget.get_config()["report_provider"]["provider_id"] == BUILTIN_PROVIDER_ID
    assert not widget.outline_editor.isReadOnly()
    widget.close()


# ═══════════════════════════════════════════════════════════════════
# Batch 3.6.6 — P0-5: DESIGN_READY confirmation control visibility
# ═══════════════════════════════════════════════════════════════════


def _make_host_workbench(qapp) -> ReportWorkbenchWidget:
    """Create a shown ReportWorkbenchWidget with PPT Master selected."""
    widget = ReportWorkbenchWidget()
    widget.set_report_provider_options((
        ReportProviderOption(
            provider_id=BUILTIN_PROVIDER_ID,
            provider_version=BUILTIN_PROVIDER_VERSION,
            display_name="内置标准生成器",
            is_builtin=True,
        ),
        ReportProviderOption(
            provider_id=PPT_MASTER_PROVIDER_ID,
            provider_version=PPT_MASTER_PROVIDER_VERSION,
            display_name="PPT Master 专业演示生成器",
        ),
    ))
    for index in range(widget.report_provider_combo.count()):
        data = widget.report_provider_combo.itemData(index) or {}
        if data.get("provider_id") == PPT_MASTER_PROVIDER_ID:
            widget.report_provider_combo.setCurrentIndex(index)
            break
    widget.show()
    qapp.processEvents()
    return widget


class TestConfirmationControlVisibility:
    """P0-5: confirmation controls must be visible, enabled, and accessible."""

    def test_scroll_area_exists_in_left_panel(self, qapp) -> None:
        """The left panel must wrap content in a QScrollArea."""
        widget = _make_host_workbench(qapp)
        splitter = widget.findChild(QSplitter)
        assert splitter is not None
        left = splitter.widget(0)
        assert isinstance(left, QScrollArea), (
            f"Left panel must be QScrollArea, got {type(left).__name__}"
        )
        assert left.widgetResizable()
        assert left.horizontalScrollBarPolicy() == Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        widget.close()

    def test_design_ready_button_visible_enabled_normal_size(self, qapp) -> None:
        """DESIGN_READY: confirmation button visible and enabled at 1200×800."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：视觉与沟通设计",
            preview_markdown="# 测试方案\n\n## 视觉与沟通设计\n...",
            next_action="confirm_design",
            next_action_label="确认设计并生成逐页计划",
            ready_for_authoring=False,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()

        btn = widget.ppt_master_confirm_btn
        assert btn is not None, "Confirmation button must exist"
        assert widget.ppt_master_group.isVisibleTo(widget), (
            "PPT Master group must be visible in widget"
        )
        assert btn.isVisibleTo(widget), (
            "Confirmation button must be visible in widget"
        )
        assert btn.isEnabled(), "Confirmation button must be enabled at DESIGN_READY"
        assert "确认设计" in btn.text(), (
            f"Button text must contain '确认设计', got: {btn.text()!r}"
        )

        geo = btn.geometry()
        assert geo.width() > 0 and geo.height() > 0, (
            f"Button geometry must be non-zero, got {geo.width()}x{geo.height()}"
        )
        widget.close()

    def test_design_ready_button_within_scroll_area_content(self, qapp) -> None:
        """Button must be laid out within the scroll area's content widget."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：视觉与沟通设计",
            preview_markdown="# 测试方案",
            next_action="confirm_design",
            next_action_label="确认设计并生成逐页计划",
        )
        qapp.processEvents()

        scroll = widget.findChild(QScrollArea)
        assert scroll is not None
        btn = widget.ppt_master_confirm_btn
        content = scroll.widget()
        assert content is not None

        # Button geometry in content coordinates
        btn_in_content = btn.mapTo(content, btn.rect().topLeft())
        content_height = content.height()

        # Button must start within positive region of content
        assert btn_in_content.y() >= 0, (
            f"Button y={btn_in_content.y()} must be >= 0 in content"
        )

        # Button bottom must not exceed content height excessively
        # (small tolerance for layout rounding)
        btn_bottom = btn_in_content.y() + btn.height()
        assert btn_bottom <= content_height + 10, (
            f"Button bottom {btn_bottom} exceeds content height {content_height}"
        )
        widget.close()

    def test_outline_ready_confirm_button_visible(self, qapp) -> None:
        """OUTLINE_READY: confirmation button must be visible."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：叙事大纲",
            preview_markdown="# 测试方案\n\n## 叙事大纲\n...",
            next_action="confirm_outline",
            next_action_label="确认大纲并生成设计方案",
            ready_for_authoring=False,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()

        assert widget.ppt_master_confirm_btn.isVisibleTo(widget)
        assert widget.ppt_master_confirm_btn.isEnabled()
        assert "确认大纲" in widget.ppt_master_confirm_btn.text()
        widget.close()

    def test_slides_ready_confirm_button_visible(self, qapp) -> None:
        """SLIDES_READY: confirmation button must be visible."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：逐页内容与图片位置",
            preview_markdown="# 测试方案\n\n## 逐页计划\n...",
            next_action="confirm_plan",
            next_action_label="确认最终逐页计划",
            ready_for_authoring=False,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()

        assert widget.ppt_master_confirm_btn.isVisibleTo(widget)
        assert widget.ppt_master_confirm_btn.isEnabled()
        assert "确认" in widget.ppt_master_confirm_btn.text()
        widget.close()

    def test_plan_confirmed_enables_generation_button(self, qapp) -> None:
        """PLAN_CONFIRMED: final generation button must be enabled."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="规划已确认，可生成 PPT Master 报告",
            preview_markdown="# 已确认方案",
            next_action="",
            next_action_label="当前阶段无需确认",
            ready_for_authoring=True,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()

        assert widget.ppt_master_confirm_btn.isVisibleTo(widget)
        assert not widget.ppt_master_confirm_btn.isEnabled(), (
            "Confirm button should be disabled at PLAN_CONFIRMED (no next action)"
        )
        assert widget.full_report_btn.isEnabled(), (
            "Full report generation button must be enabled at PLAN_CONFIRMED"
        )
        widget.close()

    def test_large_window_button_still_accessible(self, qapp) -> None:
        """Maximized/large window 1920×1080: button must still be reachable."""
        widget = _make_host_workbench(qapp)
        widget.resize(1920, 1080)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：视觉与沟通设计",
            preview_markdown="# 测试方案\n\n## 视觉与沟通设计\n配色方案...",
            next_action="confirm_design",
            next_action_label="确认设计并生成逐页计划",
            ready_for_authoring=False,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()

        btn = widget.ppt_master_confirm_btn
        assert btn.isVisibleTo(widget)
        assert btn.isEnabled()

        geo = btn.geometry()
        assert geo.width() > 0 and geo.height() > 0

        scroll = widget.findChild(QScrollArea)
        assert scroll is not None
        content = scroll.widget()
        btn_in_content = btn.mapTo(content, btn.rect().topLeft())
        assert btn_in_content.y() >= 0, (
            f"Button negative y in content: {btn_in_content.y()}"
        )
        widget.close()

    def test_status_label_visible_at_each_phase(self, qapp) -> None:
        """Status label must be visible and reflect current phase at all 3 stages."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()
        widget.set_diagnosis_record({"timestamp": "2026-08-07"})

        phases: list[tuple[str, str, str]] = [
            ("待确认：叙事大纲", "confirm_outline", "确认大纲并生成设计方案"),
            ("待确认：视觉与沟通设计", "confirm_design", "确认设计并生成逐页计划"),
            ("待确认：逐页内容与图片位置", "confirm_plan", "确认最终逐页计划"),
        ]

        for status_text, next_action, next_action_label in phases:
            widget.set_ppt_master_workflow_view(
                status_text=status_text,
                preview_markdown="# 测试方案",
                next_action=next_action,
                next_action_label=next_action_label,
                ready_for_authoring=False,
                valid_for_current_inputs=True,
            )
            qapp.processEvents()

            assert widget.ppt_master_status_label.isVisibleTo(widget), (
                f"Status label must be visible at {status_text}"
            )
            assert status_text in widget.ppt_master_status_label.text(), (
                f"Status label must contain '{status_text}', "
                f"got: {widget.ppt_master_status_label.text()!r}"
            )
            assert widget.ppt_master_confirm_btn.isVisibleTo(widget), (
                f"Confirm button must be visible at {status_text}"
            )
            assert widget.ppt_master_confirm_btn.isEnabled(), (
                f"Confirm button must be enabled at {status_text}"
            )

        widget.close()

    def test_invalidated_state_disables_confirmation(self, qapp) -> None:
        """Invalidated workflow: confirmation button must be disabled."""
        widget = _make_host_workbench(qapp)
        widget.resize(1200, 800)
        qapp.processEvents()

        widget.set_diagnosis_record({"timestamp": "2026-08-07"})
        widget.set_ppt_master_workflow_view(
            status_text="待确认：视觉与沟通设计",
            preview_markdown="# 测试方案",
            next_action="confirm_design",
            next_action_label="确认设计并生成逐页计划",
            ready_for_authoring=False,
            valid_for_current_inputs=True,
        )
        qapp.processEvents()
        assert widget.ppt_master_confirm_btn.isEnabled()

        widget.invalidate_ppt_master_workflow("需求文件已更改")
        assert not widget.ppt_master_confirm_btn.isEnabled(), (
            "Confirm button must be disabled after invalidation"
        )
        assert "已失效" in widget.ppt_master_status_label.text()
        widget.close()
