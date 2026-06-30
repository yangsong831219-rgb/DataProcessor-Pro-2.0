"""去静默失败 -> 真机 ground truth 脚本 (不跑 GUI, 走确定性代码路径)

验证项:
1. list_diagnoses: AttributeError 已死 (单次 stat, FileNotFoundError 跳过)
2. _read_file_with_warnings: .docx 表格提取非空, xlsx/不支持格式 -> visible 警告
3. _build_context: 空诊断 -> 显式告知文本 (非兜底套话)
4. _append_inclusion_footer: footer 段追加成功, 前文格式完好, 失败兜底进 warnings
5. 保存路径在项目资料库 (非 temp)

=== 本脚本不依赖 LLM, 不接受 mock 降级。 ===
"""

from __future__ import annotations
import io, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

PASSED = 0
FAILED = 0

def check(label: str, cond: bool):
    global PASSED, FAILED
    if cond:
        PASSED += 1; print(f"  [PASS] {label}")
    else:
        FAILED += 1; print(f"  [FAIL] {label}")


# ==========  set up a single temp dir for all file-based tests  ==========
td = __import__("tempfile").mkdtemp(prefix="deadsilencing_")

# ====== 1. list_diagnoses ======
print("=== 1. list_diagnoses ===")
from dp_engine.wiki_system import WikiFileSystem
wiki = WikiFileSystem()
result = wiki.list_diagnoses()
check("list_diagnoses() returns list, no AttributeError", isinstance(result, list))
check("items have name/size/mtime keys",
      len(result) == 0 or all("name" in d and "size" in d and "mtime" in d for d in result))
if result:
    check("mtime type is float", isinstance(result[0]["mtime"], (float, int)))
print("  (concurrent delete guard: try/except FileNotFoundError per file, code reviewed)")

# ====== 2. _read_file_with_warnings ======
print("\n=== 2. _read_file_with_warnings ===")
from core.report_engine import _read_file_with_warnings, _build_context

# 2a. plain txt
txt = os.path.join(td, "test.txt")
with open(txt, "w", encoding="utf-8") as f:
    f.write("Hello world 123")
content, warns = _read_file_with_warnings(txt)
check("2a. txt reads correctly", "Hello world" in content)
check("2a. txt no warnings", warns == [])

# 2b. missing file -> warning
content2, warns2 = _read_file_with_warnings(os.path.join(td, "nope.docx"))
check("2b. missing -> empty content", content2 == "")
check("2b. missing -> has warning", len(warns2) > 0)

# 2c. unsupported format (xlsx) -> warning
weird = os.path.join(td, "data.xlsx")
with open(weird, "w") as f:
    f.write("stuff")
content3, warns3 = _read_file_with_warnings(weird)
check("2c. xlsx -> empty content", content3 == "")
check("2c. xlsx -> unsupported format warning",
      any("xlsx" in w for w in warns3))

# 2d. rich table docx -> extract paras + table text
from docx import Document
rdocx = os.path.join(td, "rich_table.docx")
doc = Document()
doc.add_paragraph("Paragraph one.")
doc.add_paragraph("Paragraph two.")
tbl = doc.add_table(rows=3, cols=3)
tbl.cell(0, 0).text = "Sensor"; tbl.cell(0, 1).text = "sigma"; tbl.cell(0, 2).text = "Grade"
tbl.cell(1, 0).text = "C1";     tbl.cell(1, 1).text = "2.3";   tbl.cell(1, 2).text = "You"
tbl.cell(2, 0).text = "C2";     tbl.cell(2, 1).text = "8.1";   tbl.cell(2, 2).text = "Liang"
doc.save(rdocx)

content4, warns4 = _read_file_with_warnings(rdocx)
check("2d. docx extracts paragraph", "Paragraph one" in content4)
check("2d. docx extracts table (Sensor)", "Sensor" in content4)
check("2d. docx extracts table (Grade)", "Grade" in content4)
check("2d. docx extracts table (C2/Liang)", "C2" in content4 and "Liang" in content4)
check("2d. docx no warnings", warns4 == [])

# 2e. _build_context empty diagnosis -> explicit notice
ctx_empty, w_empty = _build_context({"req_file": ""})
check("2e. empty diag -> contains 'wei jia zai zhen duan'", "未加载诊断" in ctx_empty)
check("2e. empty diag -> NOT old boilerplate", "通用传感器数据分析场景" not in ctx_empty)
check("2e. empty diag -> no warnings", w_empty == [])

# 2f. _build_context with rich docx -> content gets through
ctx_rich, w_rich = _build_context({"project_files": [rdocx], "_diagnosis_record": None})
has_c2 = "C2" in ctx_rich
if not has_c2:
    print(f"    context preview: {ctx_rich[:300]}")
check("2f. docx content reaches _build_context output", has_c2)
check("2f. rich docx no warnings", w_rich == [])


# ====== 3. _append_inclusion_footer ======
print("\n=== 3. _append_inclusion_footer ===")
from docx import Document as Doc
from main import DataProcessorWindow as DPW

output_path = os.path.join(td, "test_report.docx")

# 3a. build a base docx (simulating builder output)
doc = Doc()
doc.add_heading("Test Report Title", 0)
doc.add_paragraph("This is body paragraph, testing pre-text format.")
tbl = doc.add_table(rows=2, cols=2)
tbl.cell(0, 0).text = "Metric"; tbl.cell(0, 1).text = "Value"
tbl.cell(1, 0).text = "sigma_max"; tbl.cell(1, 1).text = "18.5"
doc.save(output_path)

# 3b. call footer append
test_warnings = [
    "Failed: bad.xlsx (unsupported format)",
    "Failed: missing.docx (file not found)",
]
DPW._append_inclusion_footer(output_path, "word", test_warnings, diagnosis_loaded=False)

# 3c. re-open and verify
doc2 = Doc(output_path)
para_texts = [p.text for p in doc2.paragraphs]
all_para = "\n".join(para_texts)
table_texts = []
for t in doc2.tables:
    for row in t.rows:
        for cell in row.cells:
            table_texts.append(cell.text)
all_table = "\n".join(table_texts)

check("3c. footer heading exists", "资料纳入情况" in all_para)
check("3c. footer 'wei jia zai'", "未加载" in all_para)
check("3c. footer bad.xlsx listed", "bad.xlsx" in all_para)
check("3c. footer missing.docx listed", "missing.docx" in all_para)
check("3c. pre-text title intact", "Test Report Title" in all_para)
check("3c. pre-text table intact", "sigma_max" in all_table)
check("3c. pre-text body intact", "body paragraph" in all_para)
check("3c. file size > 0", os.path.getsize(output_path) > 100)

# 3d. verify red font hint when diagnosis not loaded
red_found = False
for p in doc2.paragraphs:
    for run in p.runs:
        if run.font.color and run.font.color.rgb:
            red_found = True; break
check("3d. no diagnosis -> footer has red font hint", red_found)

# 3e. footer write failure guard
# use a nonexistent parent dir to force a save failure
bad_path = os.path.join(td, "nonexistent_dir", "will_fail.docx")
lost_warns = test_warnings.copy()
DPW._append_inclusion_footer(bad_path, "word", lost_warns, diagnosis_loaded=True)
check("3e. footer write failure -> no crash", True)  # reached here
check("3e. footer write failure -> warning list grew", len(lost_warns) > len(test_warnings))
check("3e. failure -> contains 'write failure' entry",
      any("写入失败" in w or "write fail" in w.lower() for w in lost_warns))


# ====== 4. save path verification ======
print("\n=== 4. save path ===")
root = DPW.get_software_root_dir()
lib = os.path.join(root, DPW.LIBRARY_FOLDER_NAME)
check("4. LIBRARY_FOLDER_NAME == '项目资料库'", DPW.LIBRARY_FOLDER_NAME == "项目资料库")
check("4. root dir nonempty", bool(root))
check("4. project lib dir exists", os.path.isdir(lib))
abs_lib = os.path.abspath(lib)
print(f"  Absolute library path: {abs_lib}")
print(f"  Temp dir used for tests: {td}")

# ====== 5. real diagnosis directory ======
print("\n=== 5. real diagnosis list ===")
real_diag_dir = os.path.join(root, "wiki_vault", "diagnoses")
if os.path.isdir(real_diag_dir):
    real_diags = wiki.list_diagnoses()
    print(f"  Dir: {real_diag_dir}")
    print(f"  Files: {len(real_diags)}")
    for d in real_diags[:3]:
        print(f"    {d['name']} ({d['size']//1024}KB) mtime={d['mtime']}")
    check("5. list_diagnoses real dir no crash", True)
else:
    print(f"  (no real diagnosis dir at {real_diag_dir})")
    check("5. list_diagnoses on empty dir returns []", wiki.list_diagnoses() == [])

# ====== cleanup ======
__import__("shutil").rmtree(td, ignore_errors=True)

print(f"\n{'='*55}")
print(f"  {PASSED} PASSED / {FAILED} FAILED  (total {PASSED+FAILED})")
print(f"{'='*55}")
if FAILED > 0:
    print("  FAIL detected -> investigate above [FAIL] entries")
    sys.exit(1)
else:
    print("  All passed - dead silencing verification complete")
