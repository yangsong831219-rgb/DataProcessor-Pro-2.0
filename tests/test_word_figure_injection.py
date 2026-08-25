"""Word 图表 manifest 注入数量与附录布局回归测试。"""

from __future__ import annotations

import re
import zipfile
from types import SimpleNamespace

from docx import Document
from PIL import Image


class _Manifest:
    def __init__(self, figures):
        self._figures = figures
        self.count = len(figures)

    def __iter__(self):
        return iter(self._figures)


def _inline_picture_count(docx_path) -> int:
    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml")
    return xml.count(b"<wp:inline")


def test_all_eight_manifest_figures_are_injected_once(tmp_path):
    """六张 Phase B 图加两张基础图必须完整进入 Word，不能少图或重复。"""
    docx_path = tmp_path / "report.docx"
    Document().save(str(docx_path))

    figures = []
    for index in range(8):
        png_path = tmp_path / f"chart-{index + 1}.png"
        Image.new("RGB", (640, 360), color=(40 + index * 10, 100, 180)).save(png_path)
        figures.append(SimpleNamespace(
            fig_id=f"chart_{index + 1}",
            fig_no=index + 1,
            section="回归",
            title=f"图表{index + 1}",
            key_stat="ok",
            png_path=str(png_path),
            caption=f"图{index + 1} 图表{index + 1}",
        ))

    from main import _inject_figures_by_reference

    warnings: list[str] = []
    _inject_figures_by_reference(str(docx_path), _Manifest(figures), warnings)

    doc = Document(str(docx_path))
    captions = {
        p.text for p in doc.paragraphs
        if re.fullmatch(r"图\d+ 图表\d+", p.text)
    }
    assert _inline_picture_count(docx_path) == 8
    assert captions == {f"图{i} 图表{i}" for i in range(1, 9)}
    assert any("8张图" in warning for warning in warnings)


def test_manifest_audit_reports_missing_files_and_invalid_phase_b(tmp_path):
    """旧诊断记录缺图时必须 fail-loud，不能只生成一份看似成功的报告。"""
    from main import _audit_chart_manifest

    manifest = [
        SimpleNamespace(
            chart_id="data_ts_cleaning",
            module="data",
            title="清洗前后",
            produced=True,
            rel_path="missing.png",
            skip_reason="",
        ),
        SimpleNamespace(
            chart_id="phaseb_hyst_A1",
            module="phaseb",
            title="A1 迟滞回线",
            produced=False,
            rel_path="",
            skip_reason="hyst(A1: valid=0<3)",
        ),
    ]

    warnings = _audit_chart_manifest(manifest, str(tmp_path))

    assert any("诊断图文件缺失" in warning for warning in warnings)
    assert any("阶段B迟滞图没有有效数据" in warning for warning in warnings)
    assert any("重新运行阶段B" in warning for warning in warnings)


def test_uncited_figures_fall_back_to_matching_sections_not_appendix(tmp_path):
    """模型漏引图时应按图表所属模块归位，不能把全部图片堆入附件。"""
    docx_path = tmp_path / "section-placement.docx"
    doc = Document()
    doc.add_heading("数据分析与波长特征", level=1)
    doc.add_paragraph("本节分析波长时程与数据质量。")
    doc.add_heading("多源对比与一致性", level=1)
    doc.add_paragraph("本节分析不同来源的对齐结果。")
    doc.add_heading("传感器标定结果", level=1)
    doc.add_paragraph("本节分析应变标定与温度解耦。")
    doc.save(docx_path)

    figures = []
    for index, (module, title) in enumerate(
        [
            ("data_analysis", "波长差时程"),
            ("compare", "多源对比时程"),
            ("strain_calib", "A1 应变标定曲线"),
        ],
        1,
    ):
        png_path = tmp_path / f"section-{index}.png"
        Image.new("RGB", (640, 360), color=(50, 100 + index * 20, 180)).save(
            png_path
        )
        figures.append(
            SimpleNamespace(
                fig_id=f"section_{index}",
                fig_no=index,
                section=module,
                title=title,
                key_stat="",
                png_path=str(png_path),
                caption=f"图{index} {title}",
            )
        )

    from main import _inject_figures_by_reference

    warnings: list[str] = []
    _inject_figures_by_reference(str(docx_path), _Manifest(figures), warnings)

    rendered = Document(str(docx_path))
    texts = [paragraph.text for paragraph in rendered.paragraphs]
    assert "图表附录" not in texts
    for caption, next_heading in (
        ("图1 波长差时程", "多源对比与一致性"),
        ("图2 多源对比时程", "传感器标定结果"),
    ):
        assert texts.index(caption) < texts.index(next_heading)
    assert texts.index("图3 A1 应变标定曲线") > texts.index("传感器标定结果")
    assert any("按所属章节自动归位" in warning for warning in warnings)
