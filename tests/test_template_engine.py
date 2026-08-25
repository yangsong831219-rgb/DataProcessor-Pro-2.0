"""Tests for TemplateEngine — updated to current dp_engine.report_builder API."""

import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
import tempfile

from dp_engine.report_builder.models import (
    ReportSpec, Section, ContentBlock, WordReport, WordSection,
)
from dp_engine.report_builder.template_engine import TemplateEngine


class TestTemplateEngine:
    """Test suite for TemplateEngine — placeholder substitution + template filling."""

    def test_substitute_placeholders_basic(self):
        """Test basic placeholder substitution."""
        engine = TemplateEngine()

        text = 'Hello {{name}}, you have {{count}} messages.'
        values = {'name': 'Alice', 'count': 5}

        result = engine.substitute_placeholders(text, values)

        assert result == 'Hello Alice, you have 5 messages.'

    def test_substitute_placeholders_missing_key(self):
        """Test placeholder substitution with missing key keeps placeholder."""
        engine = TemplateEngine()

        text = 'Hello {{name}}, your email is {{email}}'
        values = {'name': 'Bob'}

        result = engine.substitute_placeholders(text, values)

        assert result == 'Hello Bob, your email is {{email}}'

    def test_substitute_placeholders_empty_text(self):
        """Test substitution with empty text."""
        engine = TemplateEngine()

        result = engine.substitute_placeholders('', {'key': 'value'})
        assert result == ''

    def test_fill_word_template_basic(self):
        """Test fill_word_template builds correct placeholder map from WordReport."""
        engine = TemplateEngine()

        report = WordReport(
            title='My Report',
            author='John',
            date='2024-01-01',
            sections=[
                WordSection(
                    heading='Overview',
                    content_paragraphs=['Some text content'],
                ),
            ],
        )

        replacements = engine.fill_word_template('', report)

        assert replacements['title'] == 'My Report'
        assert replacements['author'] == 'John'
        assert replacements['date'] == '2024-01-01'
        assert replacements['section_1_heading'] == 'Overview'
        assert replacements['section_1_para_1'] == 'Some text content'

    def test_substitute_multiple_placeholders(self):
        """Test substituting multiple placeholders at once."""
        engine = TemplateEngine()

        text = '{{title}} by {{author}} - {{date}}'
        values = {'title': 'Report', 'author': 'Me', 'date': '2024-12-31'}

        result = engine.substitute_placeholders(text, values)

        assert result == 'Report by Me - 2024-12-31'

    def test_substitute_no_placeholders(self):
        """Test text without placeholders remains unchanged."""
        engine = TemplateEngine()

        text = 'No placeholders here'
        values = {'key': 'value'}

        result = engine.substitute_placeholders(text, values)

        assert result == 'No placeholders here'

    def test_fill_word_template_with_docx(self):
        """Test fill_word_template with an actual docx template."""
        engine = TemplateEngine()

        report = WordReport(
            title='Template Test',
            author='Tester',
            date='2024-06-01',
            sections=[
                WordSection(
                    heading='Introduction',
                    content_paragraphs=['Welcome to the report'],
                ),
            ],
        )

        # Create a minimal docx for testing
        from docx import Document
        import io

        doc = Document()
        doc.add_paragraph('Title: {{title}}')
        doc.add_paragraph('Author: {{author}}')
        doc.add_paragraph('Date: {{date}}')

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        # Write to temp file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            f.write(buffer.read())
            temp_path = f.name

        try:
            replacements = engine.fill_word_template(temp_path, report)
            assert replacements['title'] == 'Template Test'
            assert replacements['author'] == 'Tester'
        finally:
            os.unlink(temp_path)


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
