"""Tests for WordBuilder — updated to current dp_engine.report_builder API."""

import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
import io

from docx import Document

from dp_engine.report_builder.models import (
    WordReport, WordSection, WordTable,
    ReportSpec, Section, ContentBlock,
)
from dp_engine.report_builder.word_builder import WordBuilder


class TestWordBuilder:
    """Test suite for WordBuilder — build WordReport → .docx bytes."""

    def test_build_empty_spec(self):
        """Test building a minimal WordReport."""
        report = WordReport(
            title='Test Report',
            author='Test Author',
            date='2024-01-01',
            sections=[],
        )

        builder = WordBuilder()
        result = builder.build(report)

        assert isinstance(result, bytes)
        assert len(result) > 0

        # Verify can open as docx
        doc = Document(io.BytesIO(result))
        assert doc is not None

    def test_build_with_text_content(self):
        """Test building a WordReport with text content."""
        report = WordReport(
            title='Data Analysis Report',
            author='John Doe',
            date='2024-01-15',
            sections=[
                WordSection(
                    heading='Overview',
                    content_paragraphs=['This is the overview text.'],
                ),
            ],
        )

        builder = WordBuilder()
        result = builder.build(report)

        doc = Document(io.BytesIO(result))

        # Check title
        assert any(p.text == 'Data Analysis Report' for p in doc.paragraphs)

        # Check section heading
        assert any(p.text == 'Overview' for p in doc.paragraphs)

        # Check content
        assert any('overview' in p.text.lower() for p in doc.paragraphs)

    def test_build_with_table(self):
        """Test building a WordReport with table data."""
        headers: list[str] = ['Name', 'Value', 'Status']
        rows: list[list[str | float]] = [
            ['Item 1', '100', 'Active'],
            ['Item 2', '200', 'Inactive'],
        ]

        report = WordReport(
            title='Report with Table',
            author='Jane Doe',
            date='2024-01-20',
            sections=[
                WordSection(
                    heading='Data',
                    tables=[
                        WordTable(
                            caption='Test Table',
                            headers=headers,
                            rows=rows,
                        ),
                    ],
                ),
            ],
        )

        builder = WordBuilder()
        result = builder.build(report)

        doc = Document(io.BytesIO(result))

        # Find the table
        tables = doc.tables
        assert len(tables) >= 1, f"Expected at least 1 table, got {len(tables)}"

        # First table should have data rows
        found = False
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    if 'Item 1' in cell.text or 'Name' in cell.text:
                        found = True
                        break
        assert found, "Table data not found in document"

    def test_build_with_multiple_sections(self):
        """Test building a WordReport with multiple sections."""
        report = WordReport(
            title='Multi-Section Report',
            author='Tester',
            date='2024-02-01',
            sections=[
                WordSection(
                    heading='Section 1',
                    content_paragraphs=['Content for section 1'],
                ),
                WordSection(
                    heading='Section 2',
                    content_paragraphs=['Content for section 2'],
                ),
            ],
        )

        builder = WordBuilder()
        result = builder.build(report)

        doc = Document(io.BytesIO(result))
        texts = [p.text for p in doc.paragraphs]

        assert 'Section 1' in texts
        assert 'Section 2' in texts
        assert 'Content for section 1' in texts
        assert 'Content for section 2' in texts

    def test_build_with_code_block(self):
        """Test building a WordReport with content resembling code."""
        report = WordReport(
            title='Code Report',
            author='Dev',
            date='2024-03-01',
            sections=[
                WordSection(
                    heading='Code Sample',
                    content_paragraphs=['print("Hello World")'],
                ),
            ],
        )

        builder = WordBuilder()
        result = builder.build(report)

        doc = Document(io.BytesIO(result))

        # Code-like content should be present
        assert any('print' in p.text for p in doc.paragraphs)

    def test_docx_structure(self):
        """Test that the generated docx has proper structure (ZIP header)."""
        report = WordReport(
            title='Structure Test',
            author='QA',
            date='2024-04-01',
            sections=[
                WordSection(
                    heading='Test Section',
                    content_paragraphs=['Test content'],
                ),
            ],
        )

        builder = WordBuilder()
        result = builder.build(report)

        # Should be valid ZIP (docx is a zip)
        assert result[:4] == b'PK\x03\x04'


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
