"""统一报告模板准入、标准化与缓存回归测试。"""

from __future__ import annotations

from unittest.mock import patch

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from dp_engine.report_builder.models import (
    PPTReport,
    PPTSlide,
    WordReport,
    WordSection,
)
from dp_engine.report_builder.ppt_builder import PPTBuilder
from dp_engine.report_builder.word_builder import WordBuilder
from utils.report_template_preparation import (
    TemplatePreparationResult,
    load_prepared_template_profile,
    prepare_report_template,
)


def test_word_sample_document_is_normalized_without_losing_template_style(tmp_path):
    source = tmp_path / "word_sample.docx"
    doc = Document()
    doc.sections[0].left_margin = DocxInches(1.25)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = DocxPt(14)
    doc.sections[0].header.paragraphs[0].text = "实验室页眉"
    doc.sections[0].footer.paragraphs[0].text = "受控报告"
    doc.add_heading("样例报告", level=0)
    doc.add_paragraph("这段样例正文不应进入新报告。")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "样例表格"
    doc.save(str(source))

    result = prepare_report_template(
        str(source),
        "word",
        cache_root=tmp_path / "cache",
    )

    assert result.accepted
    assert result.status == "normalized"
    assert result.usable_path is not None
    assert source.is_file()

    prepared = Document(result.usable_path)
    assert not any(paragraph.text.strip() for paragraph in prepared.paragraphs)
    assert not prepared.tables
    assert prepared.sections[0].header.paragraphs[0].text == "实验室页眉"
    assert prepared.sections[0].footer.paragraphs[0].text == "受控报告"
    assert prepared.sections[0].left_margin == DocxInches(1.25)
    assert prepared.styles["Normal"].font.name == "Arial"
    assert prepared.styles["Normal"].font.size == DocxPt(14)

    output = tmp_path / "word_output.docx"
    report = WordReport(
        title="正式报告",
        sections=[WordSection(heading="结论", content_paragraphs=["结果正常。"])],
    )
    WordBuilder().build_word_report(
        report,
        result.usable_path,
        str(output),
    )
    rendered = Document(str(output))
    assert rendered.sections[0].left_margin == DocxInches(1.25)
    assert rendered.styles["Normal"].font.name == "Arial"
    assert rendered.styles["Normal"].font.size == DocxPt(14)
    assert "这段样例正文不应进入新报告。" not in "\n".join(
        paragraph.text for paragraph in rendered.paragraphs
    )


def test_empty_word_template_is_direct_and_cache_is_reused(tmp_path):
    source = tmp_path / "blank.docx"
    Document().save(str(source))
    cache = tmp_path / "cache"

    first = prepare_report_template(str(source), "word", cache_root=cache)
    second = prepare_report_template(str(source), "word", cache_root=cache)

    assert first.status == "direct"
    assert first.accepted
    assert not first.cache_hit
    assert second.cache_hit
    assert second.usable_path == first.usable_path


def test_ppt_with_native_layouts_is_direct(tmp_path):
    source = tmp_path / "native.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "{{Report_Title}}"
    slide.placeholders[1].text = "{{Report_Author}} | {{Report_Date}}"
    presentation.save(str(source))

    result = prepare_report_template(
        str(source),
        "ppt",
        cache_root=tmp_path / "cache",
    )

    assert result.status == "direct"
    assert result.accepted
    assert result.capabilities["native_content_layout"] is True
    assert result.usable_path is not None
    profile = load_prepared_template_profile(result.usable_path)
    assert profile is not None
    assert profile["content_contract"] == "native-placeholders"


def test_ppt_without_placeholders_gets_safe_contract_and_inferred_styles(tmp_path):
    source = tmp_path / "sample_driven.pptx"
    presentation = Presentation()
    for layout in presentation.slide_layouts:
        for shape in list(layout.shapes):
            if shape.is_placeholder:
                layout.shapes._spTree.remove(shape._element)
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(0.8), Inches(3.5), Inches(1.0)
    )
    title_run = title_box.text_frame.paragraphs[0].add_run()
    title_run.text = "SAMPLE TITLE"
    title_run.font.size = Pt(40)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    body_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(4.0), Inches(7.0), Inches(1.5)
    )
    body_run = body_box.text_frame.paragraphs[0].add_run()
    body_run.text = "Sample body copy"
    body_run.font.size = Pt(22)
    body_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    content_sample = presentation.slides.add_slide(presentation.slide_layouts[0])
    content_title = content_sample.shapes.add_textbox(
        Inches(0.8), Inches(0.5), Inches(8.0), Inches(1.0)
    )
    content_title_run = content_title.text_frame.paragraphs[0].add_run()
    content_title_run.text = "CONTENT TITLE"
    content_title_run.font.size = Pt(36)
    content_title_run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    content_body = content_sample.shapes.add_textbox(
        Inches(0.8), Inches(2.5), Inches(7.0), Inches(2.0)
    )
    content_body_run = content_body.text_frame.paragraphs[0].add_run()
    content_body_run.text = "Content sample body should drive the report style."
    content_body_run.font.size = Pt(20)
    content_body_run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    presentation.save(str(source))

    result = prepare_report_template(
        str(source),
        "ppt",
        cache_root=tmp_path / "cache",
    )

    assert result.status == "normalized"
    assert result.accepted
    assert result.capabilities["native_content_layout"] is False
    assert result.usable_path is not None
    profile = load_prepared_template_profile(result.usable_path)
    assert profile is not None
    assert profile["content_contract"] == "safe-textboxes"
    assert profile["title_style"]["color_rgb"] == "111111"
    assert profile["body_style"]["color_rgb"] == "222222"
    prepared_template = Presentation(result.usable_path)
    prepared_cover_text = "\n".join(
        shape.text
        for shape in prepared_template.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "{{Report_Title}}" in prepared_cover_text
    assert "{{Report_Author}}" in prepared_cover_text
    assert "Sample body copy" not in prepared_cover_text
    prepared_meta = next(
        shape
        for shape in prepared_template.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
        and "{{Report_Author}}" in shape.text
    )
    assert str(prepared_meta.text_frame.paragraphs[0].font.color.rgb) == "FFFFFF"
    assert "  |  " not in prepared_meta.text
    prepared_slide_height = int(prepared_template.slide_height or 0)
    assert int(prepared_meta.top) >= int(prepared_slide_height * 0.42)

    output = tmp_path / "ppt_output.pptx"
    report = PPTReport(
        title="诊断报告",
        slides=[
            PPTSlide(
                slide_title="结论",
                bullet_points=["全部通道工作正常"],
            )
        ],
    )
    builder = PPTBuilder()
    builder.build_ppt_report(report, result.usable_path, str(output))
    assert not any("缺少可编辑正文占位符" in warning for warning in builder.warnings)

    rendered = Presentation(str(output))
    content_slide = rendered.slides[1]
    body_shape = next(
        shape
        for shape in content_slide.shapes
        if getattr(shape, "has_text_frame", False)
        and "全部通道工作正常" in shape.text
    )
    assert str(body_shape.text_frame.paragraphs[0].font.color.rgb) == "222222"


def test_poster_ppt_is_rejected_with_no_usable_path(tmp_path):
    source = tmp_path / "poster.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(48)
    presentation.slide_height = Inches(36)
    presentation.save(str(source))

    result = prepare_report_template(
        str(source),
        "ppt",
        cache_root=tmp_path / "cache",
    )

    assert result.status == "rejected"
    assert not result.accepted
    assert result.usable_path is None
    assert any("海报" in reason for reason in result.reasons)


def test_safe_contract_does_not_reuse_partial_body_placeholder(tmp_path):
    source = tmp_path / "partial_placeholders.pptx"
    presentation = Presentation()
    title_types = {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.VERTICAL_TITLE,
    }
    for layout in presentation.slide_layouts:
        for shape in list(layout.shapes):
            if (
                shape.is_placeholder
                and shape.placeholder_format.type in title_types
            ):
                layout.shapes._spTree.remove(shape._element)
    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover.shapes.add_textbox(
        Inches(1), Inches(1), Inches(8), Inches(1)
    ).text = "COVER"
    presentation.save(str(source))

    result = prepare_report_template(
        str(source),
        "ppt",
        cache_root=tmp_path / "cache",
    )
    assert result.status == "normalized"
    assert result.usable_path is not None

    output = tmp_path / "partial_output.pptx"
    PPTBuilder().build_ppt_report(
        PPTReport(
            title="Report",
            slides=[PPTSlide(slide_title="Result", bullet_points=["Readable body"])],
        ),
        result.usable_path,
        str(output),
    )
    rendered = Presentation(str(output))
    content = rendered.slides[1]
    body_textbox = next(
        shape
        for shape in content.shapes
        if getattr(shape, "has_text_frame", False) and "Readable body" in shape.text
    )
    assert not body_textbox.is_placeholder
    assert body_textbox.text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert not any(
        shape.placeholder_format.type in {
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
            PP_PLACEHOLDER.VERTICAL_BODY,
        }
        for shape in content.placeholders
    )


def test_workbench_stores_normalized_path_and_shows_status(qapp, tmp_path):
    from ui.report_workbench import ReportWorkbenchWidget

    source = str(tmp_path / "source.docx")
    prepared = str(tmp_path / "cache" / "prepared.docx")
    result = TemplatePreparationResult(
        report_type="word",
        status="normalized",
        source_path=source,
        usable_path=prepared,
        fingerprint="abc",
        score=88,
        reasons=("已清空样例正文。",),
        warnings=(),
        capabilities={},
        profile_path=str(tmp_path / "cache" / "template_manifest.json"),
    )
    widget = ReportWorkbenchWidget()

    with (
        patch(
            "PyQt6.QtWidgets.QFileDialog.getOpenFileName",
            return_value=(source, "Word (*.docx)"),
        ),
        patch(
            "utils.report_template_preparation.prepare_report_template",
            return_value=result,
        ),
        patch("PyQt6.QtWidgets.QMessageBox.information") as information,
    ):
        widget._on_select_template_file()

    assert widget.get_config()["template_file"] == prepared
    assert "已自动标准化" in widget._tmpl_file_label.text()
    assert widget.view_template_result_btn.isEnabled()
    information.assert_called_once()


def test_workbench_exposes_template_detection_and_normalization_entry(qapp):
    from ui.report_workbench import ReportWorkbenchWidget

    widget = ReportWorkbenchWidget()

    assert "检测/标准化" in widget.template_prepare_btn.text()
    assert "Word/PPT" in widget.template_prepare_hint.text()
    assert not widget.view_template_result_btn.isEnabled()


def test_workbench_ppt_extension_overrides_default_word_mode(qapp, tmp_path):
    from ui.report_workbench import ReportWorkbenchWidget

    source = str(tmp_path / "business_template.pptx")
    prepared = str(tmp_path / "cache" / "prepared.pptx")
    result = TemplatePreparationResult(
        report_type="ppt",
        status="direct",
        source_path=source,
        usable_path=prepared,
        fingerprint="ppt",
        score=95,
        reasons=(),
        warnings=(),
        capabilities={},
        profile_path=str(tmp_path / "cache" / "template_manifest.json"),
    )
    widget = ReportWorkbenchWidget()
    assert widget.word_radio.isChecked()

    with (
        patch(
            "PyQt6.QtWidgets.QFileDialog.getOpenFileName",
            return_value=(source, "PowerPoint (*.pptx)"),
        ),
        patch(
            "utils.report_template_preparation.prepare_report_template",
            return_value=result,
        ) as prepare,
    ):
        widget._on_select_template_file()

    prepare.assert_called_once_with(source, "ppt")
    assert widget.ppt_radio.isChecked()
    assert widget.get_config()["report_type"] == "ppt"
    assert widget.get_config()["template_file"] == prepared
    assert "模板类型: PPT" in widget._template_result_summary


def test_workbench_rejected_template_does_not_replace_current_path(qapp, tmp_path):
    from ui.report_workbench import ReportWorkbenchWidget

    source = str(tmp_path / "bad.docx")
    result = TemplatePreparationResult(
        report_type="word",
        status="rejected",
        source_path=source,
        usable_path=None,
        fingerprint="abc",
        score=0,
        reasons=("文件损坏。",),
        warnings=(),
        capabilities={},
        profile_path=None,
    )
    widget = ReportWorkbenchWidget()
    widget._template_file_path = "keep.docx"

    with (
        patch(
            "PyQt6.QtWidgets.QFileDialog.getOpenFileName",
            return_value=(source, "Word (*.docx)"),
        ),
        patch(
            "utils.report_template_preparation.prepare_report_template",
            return_value=result,
        ),
        patch("PyQt6.QtWidgets.QMessageBox.warning") as warning,
    ):
        widget._on_select_template_file()

    assert widget.get_config()["template_file"] == "keep.docx"
    warning.assert_called_once()
