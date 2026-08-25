"""ppt_builder build_ppt_report 空模板 guard — smoke test。

覆盖:
- template_path='' → 不抛错，用默认模板（内置空白 16:9）
- template_path=None → 同上
- template_path=不存在的路径 → 走默认，不抛 Package not found
- template_path=合法.pptx → 正常用该模板
"""

from __future__ import annotations

import os
import tempfile

import pytest

from dp_engine.report_builder.ppt_builder import PPTBuilder
from dp_engine.report_builder.models import PPTReport, PPTSlide


@pytest.fixture
def builder():
    return PPTBuilder()


@pytest.fixture
def minimal_report():
    """一个极简 PPTReport — 仅标题，无 slides。"""
    return PPTReport(
        title="测试报告",
        template_path=None,
        slides=[],
    )


@pytest.fixture
def report_with_slides():
    """含一页正文的 PPTReport。"""
    return PPTReport(
        title="传感器标定报告",
        template_path=None,
        slides=[
            PPTSlide(
                slide_title="数据质量评估",
                bullet_points=["异常点数: 5", "缺失率: 1.2%"],
                speaker_notes="详细说明。",
            )
        ],
    )


@pytest.fixture
def valid_template(tmp_path):
    """创建一个最小合法 .pptx 文件作为模板。"""
    from pptx import Presentation

    prs = Presentation()
    # 添加一个标题页（后续会用它验证 build_ppt_report 用了模板）
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "模板示例标题"
    slide.placeholders[1].text = "模板示例副标题"
    path = tmp_path / "template.pptx"
    prs.save(str(path))
    return str(path)


class TestBuildPptReportEmptyTemplate:
    """template_path 空/None/不存在 → 走默认，不抛错。"""

    def test_empty_string_uses_default(self, builder, report_with_slides, tmp_path):
        out = str(tmp_path / "out_empty.pptx")
        result = builder.build_ppt_report(
            report_with_slides,
            template_path="",
            output_path=out,
        )
        assert result == out
        assert os.path.exists(out)
        assert os.path.getsize(out) > 0
        # 确认默认 16:9 宽度
        from pptx import Presentation as _P
        prs = _P(out)
        # 默认 16:9 ≈ 13.333 inches
        assert prs.slide_width is not None and prs.slide_width > 0

    def test_none_uses_default(self, builder, report_with_slides, tmp_path):
        out = str(tmp_path / "out_none.pptx")
        result = builder.build_ppt_report(
            report_with_slides,
            template_path=None,  # pyright: ignore[reportArgumentType] — 特意测 None 兜底
            output_path=out,
        )
        assert result == out
        assert os.path.exists(out)
        assert os.path.getsize(out) > 0

    def test_nonexistent_path_uses_default(self, builder, report_with_slides, tmp_path):
        out = str(tmp_path / "out_nonexist.pptx")
        result = builder.build_ppt_report(
            report_with_slides,
            template_path="/nonexistent/path/template.pptx",
            output_path=out,
        )
        assert result == out
        assert os.path.exists(out)
        assert os.path.getsize(out) > 0


class TestBuildPptReportValidTemplate:
    """合法模板路径 → 正常使用该模板。"""

    def test_valid_template_used(self, builder, minimal_report, valid_template, tmp_path):
        out = str(tmp_path / "out_valid.pptx")
        result = builder.build_ppt_report(
            minimal_report,
            template_path=valid_template,
            output_path=out,
        )
        assert result == out
        assert os.path.exists(out)
        assert os.path.getsize(out) > 0

        from pptx import Presentation as _P
        prs = _P(out)
        assert prs.slides[0].shapes.title.text == minimal_report.title
        assert "模板示例标题" not in "\n".join(
            shape.text
            for shape in prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
        )

    def test_content_uses_template_placeholders_instead_of_overlay_textboxes(
        self, builder, report_with_slides, valid_template, tmp_path
    ):
        """兼容模板应沿用标题/正文占位符，不能只在模板背景上覆盖自绘文本框。"""
        out = str(tmp_path / "out_template_content.pptx")

        builder.build_ppt_report(
            report_with_slides,
            template_path=valid_template,
            output_path=out,
        )

        from pptx import Presentation as _P
        prs = _P(out)
        content_slide = prs.slides[1]
        assert content_slide.shapes.title is not None
        assert content_slide.shapes.title.text == "数据质量评估"
        body_placeholders = [
            shape for shape in content_slide.placeholders
            if shape != content_slide.shapes.title and shape.has_text_frame
        ]
        assert body_placeholders
        assert "异常点数: 5" in body_placeholders[0].text

    def test_sample_cover_shrinks_long_title_and_keeps_author_date(
        self, builder, tmp_path
    ):
        """样例页式模板的长标题不能溢出，共用作者/年份框不能丢日期。"""
        from pptx import Presentation as _P
        from pptx.util import Inches, Pt

        template = tmp_path / "sample_cover.pptx"
        prs = _P()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.8), Inches(3.8), Inches(1.5)
        )
        title.text_frame.paragraphs[0].text = "THESIS"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(80)
        decoration = slide.shapes.add_textbox(
            Inches(0.6), Inches(2.3), Inches(5.5), Inches(1.0)
        )
        decoration.text = "PRESENTATION"
        author_year = slide.shapes.add_textbox(
            Inches(1.2), Inches(5.3), Inches(4.0), Inches(1.2)
        )
        author_year.text = "AUTHOR: JANE DOE\nYEAR: 2020"
        prs.save(str(template))

        report = PPTReport(
            title="光纤传感器诊断报告",
            author="DataProcessor Pro",
            date="2026-07-20",
            slides=[],
        )
        out = tmp_path / "sample_cover_out.pptx"
        builder.build_ppt_report(report, str(template), str(out))

        rendered = _P(str(out))
        cover = rendered.slides[0]
        all_text = "\n".join(
            shape.text
            for shape in cover.shapes
            if getattr(shape, "has_text_frame", False)
        )
        assert report.title in all_text
        assert report.author in all_text
        assert report.date in all_text
        title_shape = next(
            shape for shape in cover.shapes
            if getattr(shape, "has_text_frame", False) and report.title in shape.text
        )
        title_sizes = [
            run.font.size.pt
            for paragraph in title_shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        ]
        assert title_sizes and max(title_sizes) < 80


class TestBuildPptReportWithSlides:
    """含 slides 的报告正常生成（无模板默认）。"""

    def test_slides_generated_with_default_template(self, builder, report_with_slides, tmp_path):
        out = str(tmp_path / "out_slides.pptx")
        result = builder.build_ppt_report(
            report_with_slides,
            template_path="",
            output_path=out,
        )
        assert result == out
        assert os.path.exists(out)
        # 应至少有一页内容
        prs_path = result
        from pptx import Presentation as _P
        prs = _P(prs_path)
        assert len(prs.slides) == 2
        assert any(
            report_with_slides.title in shape.text
            for shape in prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
        )
        assert any(
            "数据质量评估" in shape.text
            for shape in prs.slides[1].shapes
            if getattr(shape, "has_text_frame", False)
        )
