"""Word/PPT 模板类型与可打开性预校验。"""

from __future__ import annotations

from docx import Document
from pptx import Presentation

from utils.report_template_validation import validate_report_template


def test_pptx_template_is_not_sent_to_word_validator(tmp_path):
    template = tmp_path / "template.pptx"
    Presentation().save(str(template))

    assert validate_report_template(str(template), "ppt") is None
    assert ".docx" in (validate_report_template(str(template), "word") or "")


def test_docx_template_is_not_accepted_for_ppt(tmp_path):
    template = tmp_path / "template.docx"
    Document().save(str(template))

    assert validate_report_template(str(template), "word") is None
    assert ".pptx" in (validate_report_template(str(template), "ppt") or "")


def test_poster_sized_pptx_fails_loud(tmp_path):
    from pptx.util import Inches

    template = tmp_path / "poster.pptx"
    prs = Presentation()
    prs.slide_width = Inches(48)
    prs.slide_height = Inches(36)
    prs.save(str(template))

    error = validate_report_template(str(template), "ppt")
    assert error is not None
    assert "海报" in error or "页面尺寸" in error
