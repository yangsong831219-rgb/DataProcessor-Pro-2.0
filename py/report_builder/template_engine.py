"""Template engine for filling report placeholders with new model data."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

from .models import WordReport, WordSection, PPTReport, PPTSlide


PLACEHOLDER_RE = re.compile(r'\{\{(\w+)\}\}')


class TemplateEngine:
    """Renders report templates by replacing {{placeholder}} with content."""

    def fill_word_template(self, template_path: str, report: WordReport) -> dict[str, str]:
        """Extract placeholder map from a WordReport for template filling.

        Returns:
            Mapping of placeholder_name → replacement_text
        """
        replacements = {
            'title': report.title,
            'author': report.author,
            'date': report.date,
        }
        for i, section in enumerate(report.sections):
            replacements[f'section_{i+1}_heading'] = section.heading
            for j, para in enumerate(section.content_paragraphs):
                replacements[f'section_{i+1}_para_{j+1}'] = para
            for j, img in enumerate(section.image_anchors):
                replacements[f'section_{i+1}_img_{j+1}'] = img
        return replacements

    def fill_ppt_template(self, template_path: str, report: PPTReport) -> dict[str, str]:
        """Extract placeholder map from a PPTReport for template filling."""
        replacements = {
            'title': report.title,
            'author': report.author,
            'date': report.date,
        }
        for i, slide in enumerate(report.slides):
            replacements[f'slide_{i+1}_title'] = slide.slide_title
            replacements[f'slide_{i+1}_notes'] = slide.speaker_notes
            for j, bullet in enumerate(slide.bullet_points):
                replacements[f'slide_{i+1}_bullet_{j+1}'] = bullet
            if slide.image_anchor:
                replacements[f'slide_{i+1}_img'] = slide.image_anchor
        return replacements

    def substitute_placeholders(self, text: str, values: dict[str, Any]) -> str:
        """Substitute {{placeholder}} patterns in text."""
        return PLACEHOLDER_RE.sub(
            lambda m: str(values.get(m.group(1), m.group(0))),
            text,
        )
