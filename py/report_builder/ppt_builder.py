"""PowerPoint builder — 消费 PPTReport JSON，渲染 .pptx。

空间极简原则：每页 ≤4 bullet points，具体数据放入 Speaker Notes。
[INSERT_IMAGE: filename] 标签自动插入图片。
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from .models import PPTReport, PPTSlide


IMAGE_ANCHOR_RE = re.compile(r'\[INSERT_IMAGE:\s*([^\]]+)\]')

# 每页 Bullet 数量硬限制
MAX_BULLETS_PER_SLIDE = 4


class PPTBuilder:
    """从 PPTReport 结构体构建 .pptx 演示文稿。"""

    DEFAULT_ASSET_DIRS = ["wiki_vault/assets", "wiki_vault/pages", "project_files"]

    def __init__(self, asset_search_paths: list[str] | None = None):
        self.asset_paths = asset_search_paths or self.DEFAULT_ASSET_DIRS
        self._project_root = Path(__file__).parent.parent.parent

    # ── 主入口 ──

    def build(self, report: PPTReport) -> bytes:
        """从 PPTReport 构建 .pptx 文件。

        Args:
            report: PPTReport 结构化数据

        Returns:
            .pptx 文件二进制内容
        """
        if report.template_path and Path(report.template_path).exists():
            prs = Presentation(report.template_path)
        else:
            prs = Presentation()
            # 默认 16:9
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)

        self._add_title_slide(prs, report)

        for slide_data in report.slides:
            self._add_content_slide(prs, slide_data)

        buffer = io.BytesIO()
        prs.save(buffer)
        return buffer.getvalue()

    def build_from_json(self, json_data: dict) -> bytes:
        """从 JSON 字典构建演示文稿（便捷入口）。"""
        report = PPTReport.from_dict(json_data)
        return self.build(report)

    # ── 标题页 ──

    def _add_title_slide(self, prs: Presentation, report: PPTReport) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # 标题
        tb = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = report.title
        p.font.size = Pt(40)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        # 副标题 (作者 + 日期)
        sub = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(10), Inches(1))
        stf = sub.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        parts = []
        if report.author:
            parts.append(report.author)
        if report.date:
            parts.append(report.date)
        sp.text = '  |  '.join(parts)
        sp.font.size = Pt(20)
        sp.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        sp.alignment = PP_ALIGN.CENTER

    # ── 内容页 ──

    def _add_content_slide(self, prs: Presentation, slide_data: PPTSlide) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # --- 标题 ---
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_data.slide_title
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x18, 0x90, 0xFF)

        # 标题下划线
        line = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE
            Inches(0.8), Inches(1.25), Inches(3), Pt(3)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0x18, 0x90, 0xFF)
        line.line.fill.background()

        # --- Bullet Points ---
        bullets = slide_data.bullet_points[:MAX_BULLETS_PER_SLIDE]
        bullet_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(7), Inches(4.0))
        bf = bullet_box.text_frame
        bf.word_wrap = True

        for i, point in enumerate(bullets):
            if i == 0:
                bp = bf.paragraphs[0]
            else:
                bp = bf.add_paragraph()
            bp.text = f"• {point}"
            bp.font.size = Pt(20)
            bp.space_after = Pt(12)

        # --- 图片 (最多一张) ---
        if slide_data.image_anchor:
            img_filename = self._extract_image_filename(slide_data.image_anchor)
            if img_filename:
                self._add_slide_image(slide, img_filename)

        # --- Speaker Notes ---
        if slide_data.speaker_notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_data.speaker_notes

    def _add_slide_image(self, slide, filename: str) -> bool:
        img_path = self._resolve_image_path(filename)
        if img_path:
            try:
                slide.shapes.add_picture(
                    str(img_path),
                    Inches(8.5), Inches(1.6),
                    width=Inches(4.2),
                )
                return True
            except Exception:
                pass
        return False

    def _extract_image_filename(self, text: str) -> str | None:
        m = IMAGE_ANCHOR_RE.search(text)
        return m.group(1).strip() if m else (text.strip() or None)

    def _resolve_image_path(self, filename: str) -> Path | None:
        p = Path(filename)
        if p.is_absolute() and p.exists():
            return p
        for rel_dir in self.asset_paths:
            search_dir = self._project_root / rel_dir
            if search_dir.exists():
                for ext in ('', '.png', '.jpg', '.jpeg', '.gif', '.bmp'):
                    candidate = search_dir / f"{filename}{ext}" if not filename.endswith(ext) else search_dir / filename
                    if candidate.exists():
                        return candidate
                for child in search_dir.iterdir():
                    if child.is_dir():
                        for ext in ('', '.png', '.jpg', '.jpeg', '.gif', '.bmp'):
                            c = child / f"{filename}{ext}" if not filename.endswith(ext) else child / filename
                            if c.exists():
                                return c
        return None

    # ── 向后兼容 ──

    def build_legacy(self, spec) -> bytes:
        """从旧版 ReportSpec 构建 PPT（向后兼容）。"""
        from .models import ReportSpec
        if not isinstance(spec, ReportSpec):
            raise TypeError("Expected ReportSpec")

        report = PPTReport(
            title=spec.title,
            author=spec.author,
            date=spec.date,
            slides=[
                PPTSlide(
                    slide_title=s.title,
                    bullet_points=[
                        str(b.data)[:40] for b in s.blocks if b.type == 'text'
                    ][:MAX_BULLETS_PER_SLIDE],
                )
                for s in spec.sections
            ],
        )
        return self.build(report)
