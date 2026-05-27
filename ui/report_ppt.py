"""PPT 报告生成子页

从 main.py DataProcessorWindow 抽离的独立 QWidget。
结构与 report_word.py 镜像，仅文件类型和变量前缀不同。
"""

from __future__ import annotations

import os
from typing import List, Optional

from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication, QDialog, QFileDialog, QGroupBox, QHBoxLayout,
    QLabel, QLineEdit, QListWidget, QMessageBox, QPushButton,
    QSplitter, QTextEdit, QVBoxLayout, QWidget,
)


class PptReportWidget(QWidget):
    """PPT 报告生成页面 — 双栏仪表盘布局"""

    report_generated = pyqtSignal(str)
    report_saved = pyqtSignal(str)

    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.ppt_requirement_files: List[str] = []
        self.ppt_templates: List[str] = []
        self.ppt_project_files: List[str] = []
        self._setup_ui()

    # ── UI 构造 ──

    def _setup_ui(self) -> None:
        self.setStyleSheet("background-color: #f0f2f5;")
        main_layout = QHBoxLayout(self)
        main_layout.setContentsMargins(15, 15, 15, 15)
        main_layout.setSpacing(12)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setHandleWidth(8)
        splitter.setStyleSheet("""
            QSplitter::handle { background: linear-gradient(to bottom, #e0e0e0, #c0c0c0); }
            QSplitter::handle:hover { background: linear-gradient(to bottom, #1890ff, #40a9ff); }
        """)
        splitter.addWidget(self._build_left_panel())
        splitter.addWidget(self._build_right_panel())
        splitter.setStretchFactor(0, 40)
        splitter.setStretchFactor(1, 60)
        main_layout.addWidget(splitter)

    def _build_left_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        config_group = self._make_group("报告配置", "#1890ff")
        config_layout = QVBoxLayout(config_group)
        config_layout.setSpacing(10)

        title_row = QHBoxLayout()
        title_row.setSpacing(8)
        title_row.addWidget(self._label("报告标题:"))
        self.ppt_report_title = QLabel('未设置标题')
        self.ppt_report_title.setStyleSheet(
            'color: #333; font-weight: bold; padding: 6px 10px; '
            'background: #f5f5f5; border-radius: 4px;'
        )
        title_row.addWidget(self.ppt_report_title, stretch=1)
        title_row.addWidget(self._make_btn('编辑', '#722ed1', '#9254de', self._on_edit_title))
        config_layout.addLayout(title_row)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(8)
        btn_row.addWidget(self._make_btn('需求配置', '#1890ff', '#40a9ff', self._on_requirement_config))
        btn_row.addWidget(self._make_btn('模板操作', '#1890ff', '#40a9ff', self._on_template_operation))
        btn_row.addWidget(self._make_btn('项目文件', '#1890ff', '#40a9ff', self._on_project_files))
        config_layout.addLayout(btn_row)
        layout.addWidget(config_group)

        gen_group = self._make_group("PPT文件生成", "#52c41a")
        gen_layout = QVBoxLayout(gen_group)
        gen_layout.setSpacing(10)
        self.ppt_gen_status = QLabel('● 就绪')
        self.ppt_gen_status.setStyleSheet(
            'color: #666; font-weight: bold; padding: 8px; background: #f5f5f5; border-radius: 4px;'
        )
        gen_layout.addWidget(self.ppt_gen_status)
        gen_btn_row = QHBoxLayout()
        gen_btn_row.setSpacing(10)
        gen_btn_row.addWidget(self._make_large_btn('生成报告', '#52c41a', '#73d13d', self._on_generate))
        gen_btn_row.addWidget(self._make_large_btn('保存报告', '#1890ff', '#40a9ff', self._on_save))
        gen_btn_row.addStretch()
        gen_layout.addLayout(gen_btn_row)
        layout.addWidget(gen_group)
        layout.addStretch(1)
        return panel

    def _build_right_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        chat_group = self._make_group("AI交互", "#faad14")
        chat_layout = QVBoxLayout(chat_group)
        chat_layout.setSpacing(10)

        self.ppt_ai_chat_input = QTextEdit()
        self.ppt_ai_chat_input.setMaximumHeight(80)
        self.ppt_ai_chat_input.setPlaceholderText('输入您的指令或回复AI的提问...')
        self.ppt_ai_chat_input.setStyleSheet(
            "QTextEdit { border: 1px solid #ebebeb; border-radius: 4px; padding: 10px; font-size: 13px; }"
            "QTextEdit:focus { border-color: #faad14; }"
        )
        chat_layout.addWidget(self.ppt_ai_chat_input)

        chat_btn_row = QHBoxLayout()
        chat_btn_row.setSpacing(10)
        chat_btn_row.addWidget(self._make_btn('发送', '#1890ff', '#40a9ff', self._on_ai_send))
        chat_btn_row.addWidget(self._make_btn('清除', '#ff4d4f', '#ff7875', lambda: self.ppt_ai_chat_input.clear()))
        chat_btn_row.addStretch()
        chat_layout.addLayout(chat_btn_row)

        self.ppt_ai_result = QTextEdit()
        self.ppt_ai_result.setReadOnly(True)
        self.ppt_ai_result.setPlaceholderText('AI交互日志...\n\n提示：您可以在这里查看AI的处理过程和结果。')
        self.ppt_ai_result.setStyleSheet(
            "QTextEdit { border: 1px solid #ebebeb; border-radius: 4px; padding: 12px; "
            "font-size: 12px; line-height: 1.6; background: #fafafa; }"
        )
        chat_layout.addWidget(self.ppt_ai_result, stretch=1)
        layout.addWidget(chat_group)
        return panel

    # ── 工厂辅助 ──

    @staticmethod
    def _make_group(title: str, accent: str) -> QGroupBox:
        group = QGroupBox(title)
        group.setStyleSheet(f"""
            QGroupBox {{ border: 1px solid #e0e0e0; border-radius: 8px;
                font-weight: bold; color: {accent}; padding: 10px;
                padding-top: 22px; background: white; }}
            QGroupBox::title {{ subcontrol-origin: margin; left: 12px; padding: 0 6px; }}
        """)
        return group

    @staticmethod
    def _label(text: str) -> QLabel:
        l = QLabel(text)
        l.setStyleSheet("font-weight: bold; color: #333;")
        return l

    @staticmethod
    def _make_btn(text: str, bg: str, hover: str, handler) -> QPushButton:
        btn = QPushButton(text)
        btn.setStyleSheet(f"""
            QPushButton {{ background-color: {bg}; color: white; border: none;
                border-radius: 4px; padding: 10px 14px; font-weight: bold; }}
            QPushButton:hover {{ background-color: {hover}; }}
        """)
        btn.clicked.connect(handler)
        return btn

    @staticmethod
    def _make_large_btn(text: str, bg: str, hover: str, handler) -> QPushButton:
        btn = QPushButton(text)
        btn.setMinimumHeight(44)
        btn.setStyleSheet(f"""
            QPushButton {{ background-color: {bg}; color: white; border: none;
                border-radius: 6px; padding: 12px 24px; font-weight: bold; font-size: 14px; }}
            QPushButton:hover {{ background-color: {hover}; }}
        """)
        btn.clicked.connect(handler)
        return btn

    # ── 标题编辑 ──

    def _on_edit_title(self) -> None:
        dialog = QDialog(self)
        dialog.setWindowTitle('编辑报告标题')
        dialog.resize(500, 150)
        layout = QVBoxLayout()
        layout.addWidget(QLabel('请输入报告标题:'))
        title_input = QLineEdit()
        current = self.ppt_report_title.text()
        title_input.setText('' if current == '未设置标题' else current)
        layout.addWidget(title_input)
        btn_row = QHBoxLayout()
        ok = QPushButton('保存')
        ok.setStyleSheet('background-color: #52c41a; color: white;')
        ok.clicked.connect(lambda: self._save_title(title_input.text().strip(), dialog))
        btn_row.addWidget(ok)
        cancel = QPushButton('取消')
        cancel.clicked.connect(dialog.close)
        btn_row.addWidget(cancel)
        btn_row.addStretch()
        layout.addLayout(btn_row)
        dialog.setLayout(layout)
        dialog.exec()

    def _save_title(self, title: str, dialog: QDialog) -> None:
        if title:
            self.ppt_report_title.setText(title)
        dialog.close()

    # ── 需求配置 ──

    def _on_requirement_config(self) -> None:
        dialog = QDialog(self)
        dialog.setWindowTitle('需求配置 - PPT报告')
        dialog.resize(700, 500)
        layout = QVBoxLayout()
        layout.addLayout(self._file_dialog_btn_row(
            lambda: self._add_requirement_file(dialog),
            lambda: self._delete_requirement_file(dialog),
            lambda: self._edit_selected_file(dialog),
        ))
        self._req_list = QListWidget()
        for f in self.ppt_requirement_files:
            self._req_list.addItem(f)
        layout.addWidget(self._req_list)
        close_btn = QPushButton('关闭')
        close_btn.clicked.connect(dialog.close)
        layout.addWidget(close_btn)
        dialog.setLayout(layout)
        dialog.exec()

    def _add_requirement_file(self, dialog: QDialog) -> None:
        from PyQt6.QtCore import QSettings
        s = QSettings('DataProcessor', 'Pro')
        last_dir = s.value('ppt_req_last_dir', '')
        paths, _ = QFileDialog.getOpenFileNames(dialog, '选择需求文件', last_dir,
                                                'Text Files (*.txt *.md);;All Files (*)')
        if paths:
            s.setValue('ppt_req_last_dir', os.path.dirname(paths[0]))
            for p in paths:
                if p not in self.ppt_requirement_files:
                    self.ppt_requirement_files.append(p)
                    self._req_list.addItem(p)

    def _delete_requirement_file(self, dialog: QDialog) -> None:
        item = self._req_list.currentItem()
        if item:
            path = item.text()
            if path in self.ppt_requirement_files:
                self.ppt_requirement_files.remove(path)
                self._req_list.takeItem(self._req_list.row(item))

    # ── 模板操作 ──

    def _on_template_operation(self) -> None:
        dialog = QDialog(self)
        dialog.setWindowTitle('模板操作 - PPT报告')
        dialog.resize(700, 500)
        layout = QVBoxLayout()
        layout.addLayout(self._file_dialog_btn_row(
            lambda: self._add_template(dialog),
            lambda: self._delete_template(dialog),
            lambda: self._edit_selected_file(dialog),
        ))
        self._template_list = QListWidget()
        for t in self.ppt_templates:
            self._template_list.addItem(t)
        layout.addWidget(self._template_list)
        close_btn = QPushButton('关闭')
        close_btn.clicked.connect(dialog.close)
        layout.addWidget(close_btn)
        dialog.setLayout(layout)
        dialog.exec()

    def _add_template(self, dialog: QDialog) -> None:
        from PyQt6.QtCore import QSettings
        s = QSettings('DataProcessor', 'Pro')
        last_dir = s.value('ppt_template_last_dir', '')
        path, _ = QFileDialog.getOpenFileName(dialog, '选择PPT模板', last_dir,
                                              '文档文件 (*.txt *.csv *.xlsx *.pptx);;所有文件 (*.*)')
        if path:
            s.setValue('ppt_template_last_dir', os.path.dirname(path))
            if path not in self.ppt_templates:
                self.ppt_templates.append(path)
                self._template_list.addItem(path)

    def _delete_template(self, dialog: QDialog) -> None:
        item = self._template_list.currentItem()
        if item:
            path = item.text()
            if path in self.ppt_templates:
                self.ppt_templates.remove(path)
                self._template_list.takeItem(self._template_list.row(item))

    # ── 项目文件 ──

    def _on_project_files(self) -> None:
        dialog = QDialog(self)
        dialog.setWindowTitle('项目文件 - PPT报告')
        dialog.resize(700, 500)
        layout = QVBoxLayout()
        layout.addLayout(self._file_dialog_btn_row(
            lambda: self._add_project_file(dialog),
            lambda: self._delete_project_file(dialog),
            lambda: self._edit_selected_file(dialog),
        ))
        self._project_list = QListWidget()
        for f in self.ppt_project_files:
            self._project_list.addItem(f)
        layout.addWidget(self._project_list)
        close_btn = QPushButton('关闭')
        close_btn.clicked.connect(dialog.close)
        layout.addWidget(close_btn)
        dialog.setLayout(layout)
        dialog.exec()

    def _add_project_file(self, dialog: QDialog) -> None:
        from PyQt6.QtCore import QSettings
        s = QSettings('DataProcessor', 'Pro')
        last_dir = s.value('ppt_project_last_dir', '')
        paths, _ = QFileDialog.getOpenFileNames(
            dialog, '选择项目资料', last_dir,
            'Documents (*.doc *.docx *.pdf);;Excel (*.xlsx *.xls);;'
            'Images (*.png *.jpg *.jpeg);;Text (*.txt *.csv);;All Files (*)'
        )
        if paths:
            s.setValue('ppt_project_last_dir', os.path.dirname(paths[0]))
            for p in paths:
                if p not in self.ppt_project_files:
                    self.ppt_project_files.append(p)
                    self._project_list.addItem(p)

    def _delete_project_file(self, dialog: QDialog) -> None:
        item = self._project_list.currentItem()
        if item:
            path = item.text()
            if path in self.ppt_project_files:
                self.ppt_project_files.remove(path)
                self._project_list.takeItem(self._project_list.row(item))

    # ── 通用 ──

    @staticmethod
    def _edit_selected_file(dialog: QDialog) -> None:
        for w in dialog.findChildren(QListWidget):
            item = w.currentItem()
            if item:
                path = item.text()
                if os.path.exists(path):
                    os.startfile(path)
                else:
                    QMessageBox.warning(dialog, '警告', '文件不存在')
                return

    @staticmethod
    def _file_dialog_btn_row(on_add, on_delete, on_edit) -> QHBoxLayout:
        row = QHBoxLayout()
        for text, handler, color in [
            ('添加', on_add, '#52c41a'),
            ('删除', on_delete, '#ff4d4f'),
            ('编辑', on_edit, '#1890ff'),
        ]:
            btn = QPushButton(text)
            btn.setStyleSheet(f'background-color: {color}; color: white;')
            btn.clicked.connect(handler)
            row.addWidget(btn)
        return row

    # ── 生成报告 ──

    def _on_generate(self) -> None:
        if not self.ppt_requirement_files:
            QMessageBox.warning(self, '警告', '请先配置需求文件')
            return
        if not self.ppt_templates:
            QMessageBox.warning(self, '警告', '请先配置模板文件')
            return
        if not self.ppt_project_files:
            QMessageBox.warning(self, '警告', '请先配置项目资料')
            return

        self._set_status('状态: 正在生成PPT报告...', '#faad14')
        QApplication.processEvents()

        try:
            req_contents = []
            for rf in self.ppt_requirement_files:
                try:
                    with open(rf, 'r', encoding='utf-8') as f:
                        req_contents.append(f.read())
                except Exception as e:
                    print(f"读取需求文件失败 {rf}: {e}")

            project_contents = []
            for pf in self.ppt_project_files:
                ext = os.path.splitext(pf)[1].lower()
                try:
                    if ext in ['.txt', '.csv']:
                        with open(pf, 'r', encoding='utf-8') as f:
                            project_contents.append(f.read())
                    elif ext in ['.doc', '.docx']:
                        from docx import Document
                        doc = Document(pf)
                        project_contents.append('\n'.join(p.text for p in doc.paragraphs))
                    elif ext in ['.xlsx', '.xls']:
                        import openpyxl
                        wb = openpyxl.load_workbook(pf, data_only=True)
                        content = []
                        for sheet in wb.sheetnames:
                            ws = wb[sheet]
                            for row in ws.iter_rows(values_only=True):
                                content.append(' | '.join(str(c) if c else '' for c in row))
                        project_contents.append('\n'.join(content))
                except Exception as e:
                    print(f"读取项目资料失败 {pf}: {e}")

            title = self.ppt_report_title.text()
            if title == '未设置标题':
                title = '数据分析报告'

            prompt = (
                f"请根据以下信息生成PPT格式的数据分析报告。\n\n"
                f"报告标题: {title}\n\n"
                f"需求配置:\n{chr(10).join(req_contents)}\n\n"
                f"模板文件: {self.ppt_templates[0]}\n\n"
                f"项目资料:\n{chr(10).join(project_contents)}\n\n"
                f"请生成完整的PPT报告内容，包括幻灯片结构和大纲。"
            )

            ollama = self._find_ollama_client()
            if ollama and ollama.is_available():
                result = ollama.generate(prompt)
                if result:
                    self._set_status('状态: 已生成PPT报告', '#52c41a')
                    output = 'AI生成结果:\n' + result
                    self.ppt_ai_result.setPlainText(output)
                    self.report_generated.emit(output)
                else:
                    self._set_status('状态: PPT报告生成失败', '#ff4d4f')
            else:
                self._set_status('状态: AI不可用，请配置Ollama服务', '#faad14')

        except Exception as e:
            self._set_status(f'状态: PPT报告生成失败 - {e}', '#ff4d4f')
            import traceback
            traceback.print_exc()

    def _set_status(self, text: str, color: str) -> None:
        self.ppt_gen_status.setText(text)
        self.ppt_gen_status.setStyleSheet(
            f'color: {color}; font-weight: bold; padding: 8px; background: #f5f5f5; border-radius: 4px;'
        )

    # ── 保存报告 ──

    def _on_save(self) -> None:
        from PyQt6.QtCore import QSettings
        s = QSettings('DataProcessor', 'Pro')
        last_dir = s.value('ppt_report_save_dir', '')
        path, _ = QFileDialog.getSaveFileName(self, '保存PPT报告', last_dir,
                                              'PowerPoint Files (*.pptx)')
        if not path:
            return
        s.setValue('ppt_report_save_dir', os.path.dirname(path))
        content = self.ppt_ai_result.toPlainText()
        if content and 'AI生成结果:' in content:
            content = content.split('AI生成结果:')[1]
        try:
            from pptx import Presentation
            prs = Presentation()
            title = self.ppt_report_title.text()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title if title != '未设置标题' else '数据分析报告'
            slide.shapes.placeholders[1].text = content
            prs.save(path)
            QMessageBox.information(self, '成功', f'报告已保存至: {path}')
            self.report_saved.emit(path)
        except Exception as e:
            QMessageBox.warning(self, '警告', f'保存失败: {e}')

    # ── AI 交互 ──

    def _on_ai_send(self) -> None:
        user_input = self.ppt_ai_chat_input.toPlainText().strip()
        if not user_input:
            return
        self.ppt_ai_result.append(f'用户: {user_input}')
        self.ppt_ai_chat_input.clear()
        if any(kw in user_input for kw in ['确认', '继续', '好的', '是', 'ok', 'yes']):
            self.ppt_ai_result.append('AI: 收到确认，继续生成报告...')
        else:
            self.ppt_ai_result.append('AI: 请提供更多信息以便继续生成报告。')

    # ── 辅助 ──

    def _find_ollama_client(self):
        p = self.parent()
        while p is not None:
            if hasattr(p, 'ollama_client'):
                return p.ollama_client
            p = p.parent()
        return None
